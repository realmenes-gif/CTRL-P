"""
Ctrl-P - Memorias de calculo de estimacion (una pestaña por actividad)
=====================================================================
Mismo formato que la memoria «M.C TABIQUERIA» del proyecto:
- Encabezado con logos, datos del proyecto y período (como el formato).
- Datos de la partida (ítem, descripción, cantidad, unidad, acumulado anterior, período, acumulado).
- Un bloque por tipo (los tipos de los hitos, ej. MURO / TABIQUERÍA / TRASDOSADO; sin tipos, un bloque):
  filas = edificios (o áreas), columnas = niveles (I, II, AZOTEA…) para la cantidad de diseño, el período
  anterior, cada semana del período (cada 7 días desde el primer día, incluidos el primero y el último) y
  lo pendiente.
- Resumen para estimación.
Los totales, subtotales, pendientes, el resumen y los datos de la partida son FÓRMULAS de Excel; las
cantidades de diseño y lo ejecutado son los datos del programa.
"""

import datetime as dt
import os
import re
import sys

from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from ctrlp_core import dia_anterior, iso_a_fecha

# ------------------------------------------------------------------ encabezado (el del formato)
ENCABEZADO = {
    "proyecto": "PROYECTO: DISEÑO, CONSTRUCCIÓN, EQUIPAMIENTO Y PUESTA EN  MARCHA DEL NUEVO HOSPITAL SAN ISIDRO "
                "DE TOCOA, HONDURAS.",
    "propietario": "Propietario del Proyecto： República de Honduras, Secretaría de Estado en el Despacho de Salud "
                   "(SESAL)",
    "programa": "Número y Nombre del Programa: 2329 - Programa de Apoyo a la Red Hospitalaria de la República de "
                "Honduras",
    "coordinacion": "Coordinación: UE-PARH",
    "supervision": "Supervisión: Agencia de Apoyo a la Gestión y Supervisión (AGS VIELCA – CONCREMAT)",
    "contratista": "Contratista: Consorcio Hospital  Tocoa",
}


# ------------------------------------------------------------------ estilos (los del ejemplo)
AZUL = "0097CC"
AZUL_RES = "008BBC"
SALMON = "FBE3D6"        # tono claro del acento 2 del tema (tipos, encabezado de la partida)
SALMON_B = "FCE4D6"
CELESTE = "C1E5F5"       # tono claro del acento 1 (subtotales y totales)
BLANCO = "FFFFFF"
NUM = "#,##0.00"
CONT = '_(* #,##0.00_);_(* \\(#,##0.00\\);_(* "-"??_);_(@_)'
LETRA = "Aptos Narrow"

FINA = Side(style="thin", color="000000")
MEDIA = Side(style="medium", color="000000")


def _f(tam=12, b=False, i=False, blanca=False, nombre=LETRA):
    return Font(name=nombre, size=tam, bold=b, italic=i, color=BLANCO if blanca else "000000")


def _r(color):
    return PatternFill("solid", fgColor=color)


def _caja(izq=FINA, der=FINA, arr=FINA, aba=FINA):
    return Border(left=izq, right=der, top=arr, bottom=aba)


CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)
DER = Alignment(horizontal="right", vertical="center")

MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE",
         "NOVIEMBRE", "DICIEMBRE"]


def _mes(m, mayus=True):
    t = MESES[m - 1]
    return t if mayus else t.capitalize()


def texto_rango(d0, d1, mayus=True):
    """«16 AL 22 DE OCTUBRE» / «28 DE SEPTIEMBRE AL 4 DE OCTUBRE» (y el año si cambia)."""
    al = " AL " if mayus else " al "
    de = " DE " if mayus else " de "
    if d0.year != d1.year:
        return f"{d0.day}{de}{_mes(d0.month, mayus)}{de}{d0.year}{al}{d1.day}{de}{_mes(d1.month, mayus)}{de}{d1.year}"
    if d0.month != d1.month:
        return f"{d0.day}{de}{_mes(d0.month, mayus)}{al}{d1.day}{de}{_mes(d1.month, mayus)}"
    if d0 == d1:
        return f"{d0.day}{de}{_mes(d0.month, mayus)}"
    return f"{d0.day}{al}{d1.day}{de}{_mes(d0.month, mayus)}"


def semanas(desde, hasta):
    """Cortes de 7 dias desde el primer dia del periodo (incluidos el primero y el ultimo)."""
    d0, d1 = iso_a_fecha(desde), iso_a_fecha(hasta)
    res = []
    while d0 <= d1 and len(res) < 60:
        fin = min(d0 + dt.timedelta(days=6), d1)
        res.append((d0, fin))
        d0 = fin + dt.timedelta(days=1)
    return res


def _romano(n):
    vals = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    t = ""
    for v, s in vals:
        while n >= v:
            t += s
            n -= v
    return t


def etiqueta_nivel(nombre):
    """NIVEL 1 -> I, NIVEL 2 -> II; los demas como estan (AZOTEA, SÓTANO, TRAMO 1…)."""
    m = re.fullmatch(r"NIVEL\s+(\d+)", nombre or "")
    if m and 0 < int(m.group(1)) < 40:
        return _romano(int(m.group(1)))
    return nombre or "ÚNICA"


def _orden_nivel(nombre):
    n = (nombre or "").upper()
    if n.startswith("SÓTANO") or n.startswith("SOTANO"):
        return (0, n)
    m = re.fullmatch(r"NIVEL\s+(\d+)", n)
    if m:
        return (1, int(m.group(1)))
    if n == "AZOTEA":
        return (3, n)
    return (2, n)


def nombre_hoja(it, usados):
    base = re.sub(r"[\[\]\*\?/\\:]", "-", f"M.C {it['codigo'] or it['descripcion'][:20]}")[:31].strip() or "M.C"
    nombre, k = base, 2
    while nombre.lower() in usados:
        suf = f" ({k})"
        nombre = base[:31 - len(suf)] + suf
        k += 1
    usados.add(nombre.lower())
    return nombre


def _recurso(nombre):
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, "recursos", "memoria", nombre)


# ------------------------------------------------------------------ datos de una partida
def datos_partida(p, it, desde, hasta, cortes, solo_contractual=False):
    """Cantidades de diseño y ejecutadas por (tipo, edificio, nivel) para la memoria.
    Devuelve dict con tipos, filas (edificios/areas), columnas (niveles) y valores."""
    iid = it["id"]
    ubic = {u["nivel"]: u for u in p.ubicaciones()}
    orden_ed = {}
    for k, u in enumerate(p.ubicaciones()):
        orden_ed.setdefault(u["edificio"], k)
    hitos = {h["id"]: h for h in p.hitos_de(iid)}
    grupos = p.tipos_hito(iid)                         # [(tipo, [hitos])]
    con_tipos = any(t for t, _ in grupos)
    tipos = [t for t, _ in grupos if t] if con_tipos else [""]
    tipo_de_hito = {h["id"]: (h.get("tipo") or "") for h in hitos.values()}

    # diseño: {(tipo, nivel): cantidad}
    diseno = {}
    if con_tipos:
        for (nv, tp), c in p.cant_tipos_de(iid).items():
            if tp in tipos and nv in ubic:
                diseno[(tp, nv)] = c
    else:
        for nv, c in p.contractual_de(iid).items():
            if nv in ubic:
                diseno[("", nv)] = c
    con_diseno = bool(diseno)

    # ejecutado ponderado por dia: {(tipo, nivel, fecha): cantidad}
    ej = {}
    for r in p.con.execute("SELECT nivel, hito, fecha, SUM(cantidad) s FROM avances WHERE item=? AND fecha<=? "
                           "GROUP BY nivel, hito, fecha", (iid, hasta)):
        if r["nivel"] not in ubic:
            continue
        h = hitos.get(r["hito"] or "")
        if (r["hito"] or "") and h is None:
            continue
        pct = (h["pct"] / 100.0) if h else 1.0
        tp = tipo_de_hito.get(r["hito"] or "", "") if con_tipos else ""
        if con_tipos and tp not in tipos:
            continue
        k = (tp, r["nivel"], r["fecha"])
        ej[k] = ej.get(k, 0.0) + (r["s"] or 0.0) * pct

    # tramos de tiempo: anterior (hasta el dia antes de «desde») y cada semana
    ant_hasta = dia_anterior(desde)
    tramos_t = [(None, ant_hasta)] + [(d0.isoformat(), d1.isoformat()) for d0, d1 in cortes]

    def acum(tp, nv, hasta_):
        return sum(v for (t, n, f), v in ej.items() if t == tp and n == nv and (hasta_ is None or f <= hasta_))

    valores = {}      # (tipo, nivel) -> [anterior, sem1, sem2, ...]
    claves = set(diseno) | {(t, n) for (t, n, _f) in ej}
    for tp, nv in claves:
        if solo_contractual and con_diseno and (tp, nv) not in diseno:
            continue
        tope = diseno.get((tp, nv)) if solo_contractual else None
        previo, fila = 0.0, []
        for k, (_d0, d1) in enumerate(tramos_t):
            a = acum(tp, nv, d1)
            if tope is not None:
                a = min(a, tope)
            fila.append(a - previo)
            previo = a
        valores[(tp, nv)] = fila

    # filas (edificios) y columnas (niveles) que aparecen
    usados = {nv for (_t, nv) in diseno} | {nv for (_t, nv), f in valores.items() if any(abs(x) > 1e-9 for x in f)}
    eds = sorted({ubic[nv]["edificio"] for nv in usados}, key=lambda e: orden_ed.get(e, 0))
    nombre_ed = {u["edificio"]: u["ed_nombre"] for u in ubic.values()}
    col_de = {nv: (ubic[nv]["nv_nombre"] or "") for nv in usados}
    otras = sorted({c for c in col_de.values() if c}, key=_orden_nivel)
    if otras:
        # un area exterior sin niveles va en la primera columna (I / NIVEL 1): es practicamente lo mismo
        col_de = {nv: (c or otras[0]) for nv, c in col_de.items()}
    columnas = otras or [""]
    obs = []
    for (iid_, nv), lista in p.observaciones(desde, hasta).items():
        if iid_ == iid:
            obs += [f"{ubic[nv]['etiqueta'] if nv in ubic else ''}: {x}" for x in lista]
    return {"tipos": tipos, "con_tipos": con_tipos, "con_diseno": con_diseno, "diseno": diseno,
            "valores": valores, "edificios": [(e, nombre_ed[e]) for e in eds], "columnas": columnas,
            "col_de": col_de, "nivel_de": {(ubic[nv]["edificio"], col_de[nv]): nv for nv in usados},
            "n_tramos": len(tramos_t), "obs": obs}


# ------------------------------------------------------------------ hoja
def _celda(ws, fila, col, valor=None, font=None, fill=None, borde=None, fmt=None, al=None):
    c = ws.cell(row=fila, column=col)
    if valor is not None:
        c.value = valor
    if font is not None:
        c.font = font
    if fill is not None:
        c.fill = fill
    if borde is not None:
        c.border = borde
    if fmt is not None:
        c.number_format = fmt
    if al is not None:
        c.alignment = al
    return c


def _unir(ws, f0, c0, f1, c1, valor=None, **est):
    for f in range(f0, f1 + 1):
        for c in range(c0, c1 + 1):
            _celda(ws, f, c, **{k: v for k, v in est.items() if k != "al"})
    if (f0, c0) != (f1, c1):
        ws.merge_cells(start_row=f0, start_column=c0, end_row=f1, end_column=c1)
    return _celda(ws, f0, c0, valor, al=est.get("al"))


def _logo(ws, nombre, col, dx_px, dy_px, ancho_px, alto_px):
    ruta = _recurso(nombre)
    if not os.path.isfile(ruta):
        return
    img = XLImage(ruta)
    emu = 9525
    img.anchor = OneCellAnchor(_from=AnchorMarker(col=col, colOff=int(dx_px * emu), row=0, rowOff=int(dy_px * emu)),
                               ext=XDRPositiveSize2D(int(ancho_px * emu), int(alto_px * emu)))
    ws.add_image(img)


def _ancho_px(ws, col):
    w = ws.column_dimensions[get_column_letter(col)].width or 8.43
    return int(w * 7 + 5)


def hoja_memoria(wb, p, it, padre, desde, hasta, enc, solo_contractual=False, usados=None):
    cortes = semanas(desde, hasta)
    d = datos_partida(p, it, desde, hasta, cortes, solo_contractual)
    usados = set() if usados is None else usados
    ws = wb.create_sheet(nombre_hoja(it, usados))
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 70
    cols = d["columnas"]
    n = len(cols)
    nsem = len(cortes)
    # columnas: A tipo | diseño n | TOTAL | anterior n | semanas n c/u | pendiente n
    c_dis = 2
    sin_total = n == 1               # un solo nivel: la columna de diseño ya es el total (sin columna TOTAL)
    c_tot = c_dis + (0 if sin_total else n)
    c_ant = c_tot + 1
    c_sem = [c_ant + n * (k + 1) for k in range(nsem)]
    c_pen = c_ant + n * (nsem + 1)
    ultima_tabla = c_pen + n - 1
    W = max(ultima_tabla, 12)

    # anchos (los del ejemplo para A-L; el resto, como las semanas)
    anchos = {1: 16.9, 2: 14.3, 3: 16.7, 4: 12.4, 5: 15.2, 6: 14.4, 7: 13.7, 8: 14.4, 9: 13.0, 10: 12.9, 11: 13.2,
              12: 19.2}
    # que todo quepa: la columna A con el nombre mas largo y las numericas con el numero mas grande
    largo = max([len(e) for _, e in d["edificios"]] + [len(t or "ACTIVIDAD") for t in d["tipos"]] + [10])
    anchos[1] = max(anchos[1], largo * 1.2 + 3)
    grande = max([abs(it["cantidad"] or 0)] + [abs(v) for v in d["diseno"].values()] + [1])
    num_ancho = len(f"{grande:,.2f}") * 1.3 + 3          # total en negrita 13
    for c in range(1, W + 1):
        w = anchos.get(c, 13.5)
        if 1 < c <= ultima_tabla:
            w = max(w, num_ancho)
        if c in (2, 3, 4, 5, 6):          # el resumen usa A-F: que sus titulos quepan («EJECUTADO»)
            w = max(w, 13.6)
        ws.column_dimensions[get_column_letter(c)].width = round(w, 1)

    # ---- encabezado
    ws.row_dimensions[1].height = 91.8
    ws.row_dimensions[2].height = 45
    _unir(ws, 2, 1, 2, W, enc["proyecto"], font=_f(12, True, True, nombre="Calibri"), borde=_caja(), al=CENTRO)
    lineas = [enc["propietario"], enc["programa"], enc["coordinacion"], enc["supervision"], enc["contratista"],
              f"Período：{texto_rango(iso_a_fecha(desde), iso_a_fecha(hasta), False)} del {iso_a_fecha(hasta).year}"
              if iso_a_fecha(desde).year == iso_a_fecha(hasta).year else
              f"Período：{texto_rango(iso_a_fecha(desde), iso_a_fecha(hasta), False)}"]
    for k, t in enumerate(lineas):
        f = 3 + k
        ws.row_dimensions[f].height = 18.6
        abajo = FINA if k == len(lineas) - 1 else None
        _unir(ws, f, 1, f, W, t, font=_f(12), borde=Border(left=FINA, right=FINA, top=FINA if k == 0 else None,
                                                           bottom=abajo), al=IZQ)
    ws.row_dimensions[9].height = 18
    ws.row_dimensions[10].height = 28.8
    _unir(ws, 10, 1, 10, W, "MEMORIA DE CÁLCULO DE ESTIMACIÓN", font=_f(12, True, blanca=True), fill=_r(AZUL),
          borde=_caja(), al=CENTRO)
    ws.row_dimensions[11].height = 25.8
    _celda(ws, 11, 1, (padre or {}).get("codigo", ""), _f(12, True, nombre="Calibri"), _r(SALMON), _caja(), "@", CENTRO)
    _unir(ws, 11, 2, 11, W, (padre or {}).get("descripcion", ""), font=_f(12, True), fill=_r(SALMON_B), borde=_caja(),
          al=CENTRO)
    # fila de la partida: ITEM | DESCRIPCIÓN | CANTIDAD | UNIDAD | ACUM. ANT. | PERÍODO | ACUMULADO | OBS.
    c_cant, c_uni, c_aa, c_ep, c_ea, c_obs = W - 6, W - 5, W - 4, W - 3, W - 2, W - 1
    ws.row_dimensions[12].height = 34.2
    enc_f = dict(font=_f(8, True, blanca=True, nombre="Calibri"), fill=_r(AZUL), borde=_caja(), al=CENTRO)
    _unir(ws, 12, 1, 12, 1, "ITEM", **enc_f)
    _unir(ws, 12, 2, 12, c_cant - 1, "DESCRIPCIÓN", **enc_f)
    for c, t in ((c_cant, "CANTIDAD TOTAL"), (c_uni, "UNIDAD"), (c_aa, "ACUMULADO ANTERIOR"),
                 (c_ep, "EJECUTADO ESTE PERÍODO"), (c_ea, "EJECUTADO ACUMULADO")):
        _unir(ws, 12, c, 12, c, t, **enc_f)
    _unir(ws, 12, c_obs, 12, W, "OBSERVACIONES", **enc_f)
    ws.row_dimensions[13].height = 55.2
    dat = dict(font=_f(11, nombre="Calibri"), fill=_r(BLANCO), borde=_caja(), al=CENTRO)
    _celda(ws, 13, 1, it["codigo"], _f(11), None, _caja(), "@", Alignment(horizontal="center", vertical="center"))
    _unir(ws, 13, 2, 13, c_cant - 1, it["descripcion"], font=_f(11, nombre="Calibri"), fill=_r(BLANCO), borde=_caja(),
          al=IZQ)
    _unir(ws, 13, c_cant, 13, c_cant, it["cantidad"], fmt=NUM, **dat)
    _unir(ws, 13, c_uni, 13, c_uni, it["unidad"], **dat)
    obs = " | ".join(d["obs"]) if d["obs"] else "N/A"
    if len(obs) > 1500:
        obs = obs[:1497] + "…"
    _unir(ws, 13, c_obs, 13, W, obs, font=_f(11 if obs == "N/A" else 9, nombre="Calibri"), fill=_r(BLANCO),
          borde=_caja(), al=CENTRO)
    # alto de la fila para que la descripcion y las observaciones siempre quepan
    import math
    ancho_desc = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(2, c_cant))
    ancho_obs = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(c_obs, W + 1))
    l_desc = math.ceil(len(it["descripcion"] or "") * 1.1 / max(10, ancho_desc - 2))
    l_obs = math.ceil(len(obs) * 0.85 / max(8, ancho_obs - 2)) if obs != "N/A" else 1
    ws.row_dimensions[13].height = max(55.2, l_desc * 15 + 8, l_obs * 12.5 + 8)

    # ---- tabla
    f = 16
    ws.row_dimensions[f].height = 40.8
    th = dict(font=_f(10, True, blanca=True), fill=_r(AZUL), al=CENTRO, fmt=CONT)
    _unir(ws, f, 1, f, 1, "TIPO" if d["con_tipos"] else "UBICACIÓN",
          borde=Border(left=MEDIA, right=FINA, top=MEDIA), **th)
    _unir(ws, f, c_dis, f, c_dis + n - 1, "CANTIDAD DE DISEÑO", borde=Border(left=FINA, right=FINA, top=MEDIA), **th)
    if not sin_total:
        _unir(ws, f, c_tot, f, c_tot, "TOTAL", borde=Border(left=FINA, top=MEDIA), **th)
    _unir(ws, f, c_ant, f, c_ant + n - 1, "PERIODO ANTERIOR", borde=Border(left=MEDIA, right=FINA, top=MEDIA,
                                                                            bottom=FINA), **th)
    for k, (d0, d1) in enumerate(cortes):
        _unir(ws, f, c_sem[k], f, c_sem[k] + n - 1, "SEMANA\n" + texto_rango(d0, d1),
              borde=Border(left=MEDIA, right=FINA, top=MEDIA, bottom=FINA), **th)
    _unir(ws, f, c_pen, f, c_pen + n - 1, "PENDIENTE DE EJECUTAR",
          borde=Border(left=MEDIA, right=MEDIA, top=MEDIA, bottom=FINA), **th)
    f += 1

    grupos_ant = [c_ant] + c_sem + [c_pen]       # columnas que empiezan con borde medio
    subtotales = []

    def borde_col(c, fondo=FINA):
        izq = MEDIA if (c == 1 or c in grupos_ant) else FINA
        der = MEDIA if c == ultima_tabla else (None if (c == c_tot and not sin_total) else FINA)
        return Border(left=izq, right=der, top=FINA, bottom=fondo)

    for tp in d["tipos"]:
        # titulo del bloque: TIPO | I II | TOTAL | I II | ...
        ws.row_dimensions[f].height = 27
        tit = tp or "ACTIVIDAD"
        for c in range(1, ultima_tabla + 1):
            if c == 1:
                v = tit
            elif c == c_tot and not sin_total:
                v = "TOTAL"
            elif sin_total:
                v = etiqueta_nivel(cols[0]) if cols[0] else ""
            else:
                rel = (c - c_dis) if c < c_tot else (c - c_ant) % n
                v = etiqueta_nivel(cols[rel])
            _celda(ws, f, c, v, _f(12, True, nombre="Calibri"), _r(SALMON), borde_col(c), NUM, CENTRO)
        f += 1
        f_ini = f
        for eid, enom in d["edificios"]:
            ws.row_dimensions[f].height = 15.6
            _celda(ws, f, 1, enom, _f(12), _r(BLANCO), borde_col(1), NUM, Alignment(horizontal="center",
                                                                                     vertical="center"))
            for j, cn in enumerate(cols):
                nv = d["nivel_de"].get((eid, cn))
                dis = d["diseno"].get((tp, nv)) if nv else None
                val = d["valores"].get((tp, nv)) if nv else None
                # diseño (0 si no hay, como el ejemplo; sin datos fijos, vacio)
                _celda(ws, f, c_dis + j, (round(dis, 4) if dis is not None else 0) if d["con_diseno"] else None,
                       _f(12), _r(BLANCO), borde_col(c_dis + j), NUM)
                # anterior y semanas
                for k, c0 in enumerate([c_ant] + c_sem):
                    x = val[k] if val else 0.0
                    _celda(ws, f, c0 + j, round(x, 4) if abs(x) > 1e-9 else None, _f(12), _r(BLANCO),
                           borde_col(c0 + j), CONT, CENTRO)
                # pendiente = diseño - (anterior + semanas)
                L = get_column_letter
                suma = ",".join(f"{L(c0 + j)}{f}" for c0 in [c_ant] + c_sem)
                _celda(ws, f, c_pen + j, f"=+{L(c_dis + j)}{f}-SUM({suma})" if d["con_diseno"] else None, _f(12),
                       None, borde_col(c_pen + j), CONT)
            dis_cols = "+".join(f"{get_column_letter(c_dis + j)}{f}" for j in range(n))
            if not sin_total:
                _celda(ws, f, c_tot, f"=ROUND({dis_cols},2)" if d["con_diseno"] else None, _f(12), _r(BLANCO),
                       borde_col(c_tot), NUM, CENTRO)
            f += 1
        f_fin = f - 1
        # subtotal
        ws.row_dimensions[f].height = 19.2
        for c in range(1, ultima_tabla + 1):
            L = get_column_letter(c)
            if c == 1:
                v = "SUB TOTAL"
            elif c <= c_tot and not d["con_diseno"]:
                v = None
            elif c == c_tot and not sin_total:
                v = "=ROUND(" + "+".join(f"{get_column_letter(c_dis + j)}{f}" for j in range(n)) + ",2)"
            elif f_fin >= f_ini:
                v = f"=SUM({L}{f_ini}:{L}{f_fin})"
            else:
                v = 0
            _celda(ws, f, c, v, _f(12, True), _r(CELESTE), borde_col(c), NUM if c <= c_tot else CONT,
                   Alignment(horizontal="center" if c == 1 else "right", vertical="center"))
        subtotales.append((tit, f))
        f += 1

    # total general
    f_total = f
    ws.row_dimensions[f].height = 33.6
    for c in range(1, ultima_tabla + 1):
        L = get_column_letter(c)
        v = "TOTAL" if c == 1 else (None if (c < c_tot and not d["con_diseno"]) else
                                    "=" + "+".join(f"{L}{s}" for _, s in subtotales))
        _celda(ws, f, c, v, _f(13, True), _r(CELESTE), borde_col(c, MEDIA), NUM if c <= c_tot else CONT,
               Alignment(horizontal="center" if c == 1 else None, vertical="center", wrap_text=c == 1))

    # sin datos fijos: lo pendiente sale de la cantidad total de la partida
    L = get_column_letter
    if not d["con_diseno"]:
        _celda(ws, f_total, c_tot, f"=ROUND({L(c_cant)}13,2)")
        _celda(ws, f_total, c_pen, "=ROUND(" + L(c_tot) + str(f_total) + "-SUM(" +
               ",".join(f"{L(c0 + j)}{f_total}" for c0 in [c_ant] + c_sem for j in range(n)) + "),2)")

    # ---- datos de la partida (formulas)
    ant_total = "+".join(f"{L(c_ant + j)}{f_total}" for j in range(n))
    _unir(ws, 13, c_aa, 13, c_aa, f"=ROUND({ant_total},2)", fmt=NUM, **dat)

    # ---- resumen para estimacion
    f = f_total + 3
    ws.row_dimensions[f].height = 18
    _unir(ws, f, 1, f, 6, "RESUMEN PARA ESTIMACION", font=_f(14, True, blanca=True), fill=_r(AZUL_RES), borde=_caja(),
          al=CENTRO)
    f += 1
    ws.row_dimensions[f].height = 43.2
    for k, t in enumerate(["TIPO" if d["con_tipos"] else "ACTIVIDAD", "CANTIDAD DE DISEÑO ", "EJECUTADO ANTERIOR",
                           "EJECUTADO ESTE  PERÍODO", "EJECUTADO ACUMULADO", "PENDIENTE DE EJECUTAR"], start=1):
        _celda(ws, f, k, t, _f(11, True), _r(SALMON), _caja(), NUM, CENTRO)
    f += 1
    r_ini = f
    filas_res = subtotales if d["con_diseno"] else [(subtotales[0][0], f_total)] if subtotales else []
    for tit, s in filas_res:
        ant = "+".join(f"{L(c_ant + j)}{s}" for j in range(n))
        per = "+".join(f"{L(c0 + j)}{s}" for c0 in c_sem for j in range(n)) or "0"
        for k, v in enumerate([tit, f"=+{L(c_tot)}{s}", f"=+{ant}", f"=+{per}", f"=+C{f}+D{f}", f"=B{f}-E{f}"], 1):
            _celda(ws, f, k, v, _f(11), None, _caja(), NUM,
                   Alignment(horizontal="center", vertical="center") if k == 1 else DER)
        f += 1
    r_fin = f - 1
    for k in range(1, 7):
        v = "TOTAL" if k == 1 else f"=SUM({L(k)}{r_ini}:{L(k)}{r_fin})"
        _celda(ws, f, k, v, _f(11, True), _r(CELESTE), _caja(), NUM,
               Alignment(horizontal="center", vertical="center") if k == 1 else DER)
    f_res_total = f
    # período y acumulado de la partida
    _unir(ws, 13, c_ep, 13, c_ep, f"=D{f_res_total}", fmt=NUM, **dat)
    _unir(ws, 13, c_ea, 13, c_ea, f"={L(c_aa)}13+{L(c_ep)}13", fmt=NUM, **dat)

    f_ult = f_res_total

    # ---- logos (como el ejemplo: gobierno a la izquierda; supervision y consorcio a la derecha)
    _logo(ws, "logo_izquierda.png", 0, 10, 16, 288, 77)
    x = sum(_ancho_px(ws, c) for c in range(1, W + 1))
    x_cons = x - 119 - 8
    x_sup = x_cons - 205 - 18
    for nombre, xx, yy, an, al_ in (("logo_supervision.png", x_sup, 13, 205, 56),
                                    ("logo_contratista.png", x_cons, 5, 119, 73)):
        col, acum_ = 0, 0
        while col < W and acum_ + _ancho_px(ws, col + 1) <= xx:
            acum_ += _ancho_px(ws, col + 1)
            col += 1
        _logo(ws, nombre, col, xx - acum_, yy, an, al_)

    # ---- impresion (como el ejemplo)
    ws.print_area = f"A1:{get_column_letter(W)}{f_ult}"
    ws.print_title_rows = "16:16"
    ws.page_setup.orientation = "portrait" if W <= 14 else "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left, ws.page_margins.right = 0.3937, 0.1181
    ws.page_margins.top, ws.page_margins.bottom = 0.7874, 0.1969
    ws.page_margins.header = ws.page_margins.footer = 0.315
    ws.sheet_view.view = "pageBreakPreview"
    return ws


def compactar_imagenes(ruta):
    """Excel guarda los logos una vez por pestaña: se deja una sola copia de cada uno y todas las pestañas
    la usan (el archivo pesa mucho menos y se ve igual)."""
    import hashlib
    import zipfile
    tmp = ruta + ".tmp"
    with zipfile.ZipFile(ruta) as z:
        nombres = z.namelist()
        media = [n for n in nombres if n.startswith("xl/media/")]
        unico, cambio = {}, {}
        for n in media:
            h = hashlib.sha256(z.read(n)).hexdigest()
            if h in unico:
                cambio[n.split("/")[-1]] = unico[h].split("/")[-1]
            else:
                unico[h] = n
        if not cambio:
            return
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
            for n in nombres:
                if n.startswith("xl/media/") and n.split("/")[-1] in cambio:
                    continue
                datos = z.read(n)
                if n.startswith("xl/drawings/_rels/"):
                    t = datos.decode("utf-8")
                    for viejo, nuevo in cambio.items():
                        t = t.replace(f"/media/{viejo}\"", f"/media/{nuevo}\"")
                    datos = t.encode("utf-8")
                out.writestr(z.getinfo(n), datos)
    os.replace(tmp, ruta)


def padre_de(p, it, por_id):
    return por_id.get(it.get("padre")) if it.get("padre") else None


def agregar_memorias(wb, p, ids, desde, hasta, solo_contractual=False, progreso=None, f0=0.0, f1=1.0):
    """Agrega una pestaña de memoria por cada partida (tipo «P») de ids, en el orden del presupuesto."""
    enc = ENCABEZADO
    items = p.items()
    por_id = {i["id"]: i for i in items}
    lista = [i for i in items if i["tipo"] == "P" and i["id"] in ids]
    usados = {ws.title.lower() for ws in wb.worksheets}
    for k, it in enumerate(lista):
        if progreso:
            progreso(f0 + (f1 - f0) * k / max(1, len(lista)), f"Memoria {it['codigo']} ({k + 1} de {len(lista)})…")
        hoja_memoria(wb, p, it, padre_de(p, it, por_id), desde, hasta, enc, solo_contractual, usados)
    return len(lista)

"""
Ctrl-P - formatos propios de memoria de calculo por actividad
=============================================================
Las actividades que tienen formato en «formatos_de_memorias.xlsm» (se reconocen por el codigo del item)
salen con ese orden de filas y columnas. El calculo es el mismo de siempre (lo ejecutado ponderado por hitos,
el tope contractual, el periodo anterior y cada semana del periodo o «ESTE PERÍODO» en el corte de
estimacion); solo cambia como se presenta.
En todas: cantidad de diseño, anterior, este periodo (semanas), acumulado y pendiente; mas las columnas
especiales de cada formato (% de avance, costo unitario, ejecutado por hito…).
Encabezado con logos, sin firmas.
"""

import re

from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties

import ctrlp_memoria as M
from ctrlp_core import dia_anterior, iso_a_fecha

DDEBF7 = "DDEBF7"      # totales de la tabla de hitos (como el ejemplo)
FFF2CC = "FFF2CC"      # unidades base del hito 1
PCT = "0.00%"
PCT0 = "0%"


# ------------------------------------------------------------------ plantillas (por codigo de item)
def _cod(c):
    """«04.05», «4.05» y «4.05.» son el mismo codigo."""
    partes = [x for x in str(c or "").strip().strip(".").split(".") if x.strip() != ""]
    try:
        return tuple(int(x) for x in partes)
    except ValueError:
        return tuple(x.strip().upper() for x in partes)


NOMBRES = {"dis": "CANTIDAD DE DISEÑO ({u})", "ant": "ACUMULADO ANTERIOR", "per": "EJECUTADO ESTE PERÍODO",
           "acu": "EJECUTADO ACUMULADO", "pen": "PENDIENTE DE EJECUTAR"}


def _p(**kw):
    d = {"filas": "edificio", "etq": [("EDIFICIO", "ed")], "split": (), "total_dis": True, "nombres": {},
         "extras": (), "titulo": "", "hitos_per": False, "etq_nivel": None}
    d.update(kw)
    return d


_LISTA_PCT = ("pct",)
PLANTILLAS = {
    # --- por espacio / edificio (una fila por edificio)
    "4.05": _p(etq=[("ESPACIO", "ed")], extras=_LISTA_PCT, nombres={"pct": "% AVANCE"}),
    "4.06": _p(etq=[("ESPACIO", "ed")], extras=_LISTA_PCT,
               nombres={"dis": "ÁREA DE DISEÑO ({u})", "pct": "% ACUMULADO"}),
    "9.02.01": _p(etq=[("No", "num"), ("EDIFICIO", "ed")], extras=_LISTA_PCT, nombres={"pct": "% ACUMULADO"}),
    "20.01.05": _p(etq=[("ZONA", "ubic")], filas="ubicacion_plana", extras=_LISTA_PCT, nombres={"pct": "% AVANCE"}),
    "20.01.07": _p(etq=[("ZONA", "ubic")], filas="ubicacion_plana", extras=_LISTA_PCT, nombres={"pct": "% AVANCE"}),
    "22.01.04": _p(etq=[("TIPO", "tipo"), ("EDIFICIO/NIVEL", "ubic")], filas="ubicacion_plana", extras=_LISTA_PCT,
                   nombres={"pct": "% AVANCE"}),
    # --- edificio con sus niveles debajo
    "6.01": _p(titulo="MURO DE BLOQUE DE CONCRETO DE 3000 PSI", filas="ubicacion",
               etq=[("ESPACIO", "ed"), ("NIVEL", "nivel_num")],
               nombres={"dis": "DISEÑO {u}", "ant": "EJECUTADO ACUMULADO ANTERIOR"}),
    # --- niveles en sub-columnas
    "5.02": _p(etq=[("MODULO", "ed")], split=("dis", "ant", "per", "acu", "pen", "pct"),
               extras=("pct",), nombres={"pct": "PORCENTAJE ACUMULADO", "pen": "PENDIENTE DE COBRO"},
               etq_nivel="entrepiso"),
    "6.03": _p(etq=[("TIPO", "ed")], split=("dis", "ant", "per", "acu", "pen"),
               nombres={"dis": "CANTIDAD DE DISEÑO", "ant": "EJECUTADO ANTERIOR"}, etq_nivel="romano"),
    "7.02.01": _p(etq=[("TIPO", "ed")], split=("dis", "ant", "per", "acu", "pen"),
                  nombres={"dis": "CANTIDAD DE DISEÑO", "ant": "EJECUTADO ANTERIOR"}, etq_nivel="romano"),
    "7.01.01.01": _p(etq=[("No", "num"), ("EDIFICIO", "ed")], split=("dis", "per"),
                     nombres={"dis": "CANTIDAD DE DISEÑO"}),
    "7.01.01.02": _p(etq=[("No", "num"), ("EDIFICIO", "ed")], split=("dis", "per"),
                     nombres={"dis": "CANTIDAD DE DISEÑO"}),
    "18.01.02.01": _p(titulo="SISTEMA DE DRENAJE SANITARIO Ø 2\"", etq=[("No", "num"), ("EDIFICIO", "ed"),
                                                                       ("UNIDAD", "unidad")],
                      split=("dis",), nombres={"dis": "CANTIDAD DISEÑO ({u})", "per": "EJECUTADO"}, hitos_per=True),
    "18.01.02.03": _p(titulo="SISTEMA DE DRENAJE SANITARIO Ø 4\"", etq=[("No", "num"), ("EDIFICIO", "ed"),
                                                                       ("UNIDAD", "unidad")],
                      split=("dis",), nombres={"dis": "CANTIDAD DISEÑO ({u})", "per": "EJECUTADO"}, hitos_per=True),
    "18.01.02.04": _p(titulo="SISTEMA DE DRENAJE SANITARIO Ø 6\"", etq=[("No", "num"), ("EDIFICIO", "ed"),
                                                                       ("UNIDAD", "unidad")],
                      split=("dis",), nombres={"dis": "CANTIDAD DISEÑO ({u})", "per": "EJECUTADO"}, hitos_per=True),
    "18.01.01.02": _p(titulo="SISTEMA DE DRENAJE SANITARIO", etq=[("No", "num"), ("EDIFICIO", "ed"),
                                                                 ("DIAMETRO", "tipo"), ("UNIDAD", "unidad")],
                      split=("dis",), nombres={"dis": "CANTIDAD DISEÑO ({u})", "per": "EJECUTADO"}, hitos_per=True,
                      tipos_en_filas=True),
    # --- secciones por nivel, con los tipos / equipos de cada edificio
    "20.05.07": _p(filas="seccion", etq=[("EDIFICIO", "ed"), ("TIPO DE LAMPARA", "sub")],
                   nombres={"dis": "CANTIDAD DE DISEÑO ({u})"}, extras=_LISTA_PCT),
    "20.05.08": _p(filas="seccion", etq=[("EDIFICIO", "ed"), ("TOMACORRIENTES", "sub")],
                   nombres={"dis": "CANTIDAD DE DISEÑO ({u})"}, extras=_LISTA_PCT),
    "21.17.01": _p(filas="seccion", etq=[("BLOQUE", "ed"), ("CANTIDAD DE DISEÑO BLOQUE ({u})", "dis_ed"),
                                         ("DUCTOS DE EQUIPO", "sub")],
                   nombres={"dis": "CANTIDAD DE DISEÑO ({u})"}, extras=_LISTA_PCT),
    # --- sectores y tramos
    "17.02.02": _p(titulo="VEREDA DE CONCRETO", filas="ubicacion", etq=[("SECTOR", "ed"), ("No", "nivel_num")],
                   nombres={"dis": "CANTIDAD DISEÑO ({u})", "per": "EJECUTADO"}),
    # --- fachadas (tramos) con costo
    "6.05": _p(filas="tramo", etq=[("N°", "num"), ("DESCRIPCIÓN", "tramo"), ("SISTEMA", "nivel_rom")],
               nombres={"dis": "ÁREA ({u})"}, extras=("pu", "costo", "monto_per", "monto_pen")),
}
PLANTILLAS = {_cod(k): v for k, v in PLANTILLAS.items()}

# codigos con formato de pago por hitos en el ejemplo (las demas partidas con hitos de equipos tambien)
HITOS_EJEMPLO = {_cod(c) for c in (
    "11.01.02.01 11.01.02.03 11.01.02.05 11.01.02.08 11.01.02.100 11.01.02.102 20.02.07 20.02.10 20.02.11 "
    "20.02.12 20.02.13 20.02.14 20.02.15 20.02.16 20.02.17 20.02.18 20.02.19 20.02.21 20.02.22 20.09.09 21.20.01 "
    "21.20.02 21.20.03 21.20.04 25.02.01 25.03.01 25.03.02 25.03.03 25.03.05 25.04.01 25.04.02 25.06.01 25.06.02 "
    "25.06.03 25.06.04 25.06.05 25.06.06 25.06.07 25.06.08 25.06.09 25.06.12 25.06.13 28.01").split()}

# hitos del ejemplo: los usan las partidas de la lista que todavia no tienen hitos cargados en Datos fijos
HITOS_ESTANDAR = [{"id": f"_estandar{k}", "nombre": n, "pct": pc, "tipo": ""} for k, (n, pc) in enumerate((
    ("Despacho en puerto de origen", 40.0), ("Equipos arribados al país y en bodega", 20.0),
    ("Instalación y puesta en obra", 25.0), ("Fase de pruebas y aceptación", 15.0)))]

REQUISITOS = {
    "despacho en puerto de origen": "• Orden de compra  •  Factura  •  Certificados de origen  •  Evidencias "
                                    "fotográficas  •  Pruebas de aceptación en fábrica (FAT)  •  Guía de embarque  •  "
                                    "Póliza de seguro",
    "equipos arribados al país y en bodega": "• Inspección física  •  Acta de inspección del equipo  •  "
                                             "Verificación de condiciones de almacenamiento",
    "instalación y puesta en obra": "• Montaje físico  •  Conexionado electromecánico  •  Obra civil "
                                    "correspondiente",
    "fase de pruebas y aceptación": "• Pruebas de funcionamiento  •  Acta de recepción provisional  •  Entrega de "
                                    "protocolos de prueba  •  Acta debidamente firmada",
}


def _norm(t):
    return re.sub(r"\s+", " ", (t or "").strip().lower())


def plantilla_de(p, it):
    """Formato propio de la partida: ('hitos', None), ('tabla', plantilla) o None (formato general)."""
    c = _cod(it.get("codigo"))
    hitos = p.hitos_de(it["id"])
    sin_tipos = hitos and not any(h.get("tipo") for h in hitos)
    if sin_tipos and (c in HITOS_EJEMPLO or any(_norm(h["nombre"]).startswith("despacho") for h in hitos)):
        return ("hitos", None)
    if not hitos and c in HITOS_EJEMPLO and not p.con.execute(
            "SELECT COUNT(*) FROM avances WHERE item=? AND cantidad<>0", (it["id"],)).fetchone()[0]:
        return ("hitos", None)          # todavia sin hitos ni avance: con los hitos del ejemplo
    if c in PLANTILLAS:
        return ("tabla", PLANTILLAS[c])
    return None


# ------------------------------------------------------------------ datos
class Datos:
    """Diseño y ejecutado de una partida por (tipo, nivel, tramo, equipo) y por hito, con la misma mecanica de
    la memoria general (ponderado por el % de cada hito; con tope contractual si se pide)."""

    def __init__(self, p, it, desde, hasta, cortes, solo_contractual=False):
        iid = it["id"]
        self.it = it
        ubs = p.ubicaciones()
        self.ubic = {u["nivel"]: u for u in ubs}
        self.orden_nv = {u["nivel"]: k for k, u in enumerate(ubs)}
        self.orden_ed = {}
        for k, u in enumerate(ubs):
            self.orden_ed.setdefault(u["edificio"], k)
        self.nombre_ed = {u["edificio"]: u["ed_nombre"] for u in ubs}
        self.hitos = p.hitos_de(iid)
        self.hito = {h["id"]: h for h in self.hitos}
        grupos = p.tipos_hito(iid)
        self.con_tipos = any(t for t, _ in grupos)
        self.tipos = [t for t, _ in grupos if t] if self.con_tipos else [""]
        tipo_de_hito = {h["id"]: (h.get("tipo") or "") for h in self.hitos}
        self.tramos = {t["id"]: t for t in p.tramos_de(iid) if t["nivel"] in self.ubic}
        self.equipos = {e["id"]: e for e in p.equipos_de(iid) if e["nivel"] in self.ubic}
        self.orden_tr = {t: k for k, t in enumerate(self.tramos)}
        self.orden_eq = {e: k for k, e in enumerate(self.equipos)}

        # diseño por clave (tipo, nivel, tramo, equipo)
        dis = {}
        if self.con_tipos:
            for (nv, tp), c in p.cant_tipos_de(iid).items():
                if tp in self.tipos and nv in self.ubic:
                    dis[(tp, nv, "", "")] = c
        else:
            for nv, c in p.contractual_de(iid).items():
                if nv in self.ubic:
                    dis[("", nv, "", "")] = c
            # subdivisiones con cantidad propia (equipos o tramos): la del nivel va repartida en ellas
            for e in self.equipos.values():
                if e["contractual"] is not None:
                    dis[("", e["nivel"], "", e["id"])] = e["contractual"]
                    dis.pop(("", e["nivel"], "", ""), None)
            for t in self.tramos.values():
                if t["contractual"] is not None:
                    eq = t.get("equipo") or ""
                    dis[("", t["nivel"], t["id"], eq)] = t["contractual"]
                    dis.pop(("", t["nivel"], "", ""), None)
                    dis.pop(("", t["nivel"], "", eq), None)
        self.dis = dis
        self.con_diseno = bool(dis)

        # ejecutado por dia: {(tipo, nivel, tramo, equipo, hito, fecha): (ponderado, unidades)}
        ej = {}
        for r in p.con.execute("SELECT nivel, hito, tramo, equipo, fecha, SUM(cantidad) s FROM avances WHERE item=? AND "
                               "fecha<=? GROUP BY nivel, hito, tramo, equipo, fecha", (iid, hasta)):
            if r["nivel"] not in self.ubic:
                continue
            hid = r["hito"] or ""
            h = self.hito.get(hid)
            if hid and h is None:
                continue
            pct = (h["pct"] / 100.0) if h else 1.0
            tp = tipo_de_hito.get(hid, "") if self.con_tipos else ""
            if self.con_tipos and tp not in self.tipos:
                continue
            tr = r["tramo"] or ""
            tr = tr if tr in self.tramos else ""
            eq = r["equipo"] or ""
            eq = eq if eq in self.equipos else ""
            k = (tp, r["nivel"], tr, eq, hid, r["fecha"])
            a, u = ej.get(k, (0.0, 0.0))
            ej[k] = (a + (r["s"] or 0.0) * pct, u + (r["s"] or 0.0))

        ant_hasta = dia_anterior(desde)
        self.cortes = cortes
        tramos_t = [ant_hasta] + [d1.isoformat() for _d0, d1 in cortes]
        self.n = len(tramos_t)

        def por_tramo(filas):
            """[(fecha, valor)] -> [anterior, sem1, ...]"""
            res, previo = [], 0.0
            for d1 in tramos_t:
                a = sum(v for f, v in filas if f <= d1)
                res.append(a - previo)
                previo = a
            return res

        # serie por clave base (tipo, nivel, tramo, equipo) con tope y su reparto por hito
        claves = {k[:4] for k in ej} | set(dis)
        self.serie_base = {}
        self.serie_hito = {}      # (clave base, hito) -> serie ponderada
        self.unid_hito = {}       # hito -> serie de unidades (sin ponderar)
        for kb in claves:
            if solo_contractual and self.con_diseno and kb not in dis:
                continue
            filas = [(k[5], v[0]) for k, v in ej.items() if k[:4] == kb]
            tope = dis.get(kb) if solo_contractual else None
            previo, serie = 0.0, []
            for d1 in tramos_t:
                a = sum(v for f, v in filas if f <= d1)
                if tope is not None:
                    a = min(a, tope)
                serie.append(a - previo)
                previo = a
            self.serie_base[kb] = serie
            # por hito: misma proporcion que lo que queda tras el tope
            for hid in {k[4] for k in ej if k[:4] == kb}:
                sh = por_tramo([(k[5], v[0]) for k, v in ej.items() if k[:4] == kb and k[4] == hid])
                tot = por_tramo(filas)
                self.serie_hito[(kb, hid)] = [x * (s / t) if abs(t) > 1e-12 else 0.0
                                              for x, s, t in zip(sh, serie, tot)]
        for hid in {k[4] for k in ej}:
            self.unid_hito[hid] = por_tramo([(k[5], v[1]) for k, v in ej.items() if k[4] == hid])
        self.claves = set(self.serie_base) | set(dis)

    # --- consultas
    def _ok(self, kb, f):
        tp, nv, tr, eq = kb
        u = self.ubic.get(nv)
        if u is None:
            return False
        for campo, val in f.items():
            if val is None:
                continue
            if campo == "tp" and tp != val:
                return False
            if campo == "nv" and nv != val:
                return False
            if campo == "nvs" and nv not in val:
                return False
            if campo == "ed" and u["edificio"] != val:
                return False
            if campo == "tr" and tr != val:
                return False
            if campo == "eq" and eq != val:
                return False
        return True

    def diseno(self, **f):
        ks = [k for k in self.dis if self._ok(k, f)]
        return sum(self.dis[k] for k in ks) if ks else None

    def serie(self, hito=None, **f):
        res = [0.0] * self.n
        for kb, s in self.serie_base.items():
            if not self._ok(kb, f):
                continue
            if hito is not None:
                s = self.serie_hito.get((kb, hito), [0.0] * self.n)
            res = [a + b for a, b in zip(res, s)]
        return res

    def hay(self, **f):
        """Hay diseño o algo ejecutado con ese filtro."""
        return any(self._ok(k, f) and (k in self.dis or any(abs(x) > 1e-9 for x in self.serie_base.get(k, [])))
                   for k in self.claves)

    def niveles_usados(self, **f):
        return sorted({k[1] for k in self.claves if self._ok(k, f) and self.hay(nv=k[1], **f)},
                      key=lambda nv: self.orden_nv.get(nv, 0))

    def edificios_usados(self, **f):
        return sorted({self.ubic[nv]["edificio"] for nv in self.niveles_usados(**f)},
                      key=lambda e: self.orden_ed.get(e, 0))


# ------------------------------------------------------------------ etiquetas de nivel
def _etq_nivel(nombre, modo):
    n = (nombre or "").upper()
    if modo == "romano":
        return M.etiqueta_nivel(n) if n else "ÚNICA"
    if modo == "entrepiso":
        if n == "AZOTEA":
            return "AZOTEA"
        return "ENTREPISO" if n else "ÚNICA"
    return n or "ÚNICA"


def _num_nivel(nombre):
    m = re.fullmatch(r"NIVEL\s+(\d+)", (nombre or "").upper())
    return int(m.group(1)) if m else (nombre or "—")


# ------------------------------------------------------------------ encabezado (igual que la memoria general)
def _ancestros(p, it):
    por_id = {i["id"]: i for i in p.items()}
    res, x = [], por_id.get(it.get("padre"))
    while x is not None and len(res) < 3:
        res.insert(0, x)
        x = por_id.get(x.get("padre"))
    return res


def _encabezado(ws, p, it, W, desde, hasta, enc, anchos_fijos=None):
    """Filas 1-10 como la memoria general, los capitulos del item y la fila del item.
    Devuelve (fila del item, columnas {cant, uni, aa, ep, ea, obs}, estilo de dato)."""
    ws.row_dimensions[1].height = 91.8
    ws.row_dimensions[2].height = 45
    M._unir(ws, 2, 1, 2, W, enc["proyecto"], font=M._f(12, True, True, nombre="Calibri"), borde=M._caja(), al=M.CENTRO)
    d0, d1 = iso_a_fecha(desde), iso_a_fecha(hasta)
    lineas = [enc["propietario"], enc["programa"], enc["coordinacion"], enc["supervision"], enc["contratista"],
              f"Período：{M.texto_rango(d0, d1, False)} del {d1.year}" if d0.year == d1.year else
              f"Período：{M.texto_rango(d0, d1, False)}"]
    for k, t in enumerate(lineas):
        f = 3 + k
        ws.row_dimensions[f].height = 18.6
        M._unir(ws, f, 1, f, W, t, font=M._f(12), borde=Border(left=M.FINA, right=M.FINA,
                                                               top=M.FINA if k == 0 else None,
                                                               bottom=M.FINA if k == len(lineas) - 1 else None),
                al=M.IZQ)
    ws.row_dimensions[9].height = 18
    ws.row_dimensions[10].height = 28.8
    M._unir(ws, 10, 1, 10, W, "MEMORIA DE CÁLCULO DE ESTIMACIÓN", font=M._f(12, True, blanca=True),
            fill=M._r(M.AZUL), borde=M._caja(), al=M.CENTRO)
    f = 11
    import math
    ancho_cap = sum(ws.column_dimensions[L(x)].width or 9 for x in range(2, W + 1))
    for a in _ancestros(p, it):
        lineas_cap = math.ceil(len(a.get("descripcion") or "") * 1.25 / max(20, ancho_cap - 4))
        ws.row_dimensions[f].height = max(25.8, lineas_cap * 17 + 6)
        M._celda(ws, f, 1, a.get("codigo", ""), M._f(12, True, nombre="Calibri"), M._r(M.SALMON), M._caja(), "@",
                 M.CENTRO)
        M._unir(ws, f, 2, f, W, a.get("descripcion", ""), font=M._f(12, True), fill=M._r(M.SALMON_B),
                borde=M._caja(), al=M.CENTRO)
        f += 1
    c = {"cant": W - 6, "uni": W - 5, "aa": W - 4, "ep": W - 3, "ea": W - 2, "obs": W - 1}
    ws.row_dimensions[f].height = 34.2
    enc_f = dict(font=M._f(8, True, blanca=True, nombre="Calibri"), fill=M._r(M.AZUL), borde=M._caja(), al=M.CENTRO)
    M._unir(ws, f, 1, f, 1, "ITEM", **enc_f)
    M._unir(ws, f, 2, f, c["cant"] - 1, "DESCRIPCIÓN", **enc_f)
    for col, t in ((c["cant"], "CANTIDAD TOTAL"), (c["uni"], "UNIDAD"), (c["aa"], "ACUMULADO ANTERIOR"),
                   (c["ep"], "EJECUTADO ESTE PERÍODO"), (c["ea"], "EJECUTADO ACUMULADO")):
        M._unir(ws, f, col, f, col, t, **enc_f)
    M._unir(ws, f, c["obs"], f, W, "OBSERVACIONES", **enc_f)
    f += 1
    dat = dict(font=M._f(11, nombre="Calibri"), fill=M._r(M.BLANCO), borde=M._caja(), al=M.CENTRO)
    M._celda(ws, f, 1, it["codigo"], M._f(11), None, M._caja(), "@", Alignment(horizontal="center", vertical="center"))
    M._unir(ws, f, 2, f, c["cant"] - 1, it["descripcion"], font=M._f(11, nombre="Calibri"), fill=M._r(M.BLANCO),
            borde=M._caja(), al=M.IZQ)
    M._unir(ws, f, c["cant"], f, c["cant"], it["cantidad"], fmt=M.NUM, **dat)
    M._unir(ws, f, c["uni"], f, c["uni"], it["unidad"], **dat)
    M._unir(ws, f, c["obs"], f, W, None, font=M._f(11, nombre="Calibri"), fill=M._r(M.BLANCO), borde=M._caja(),
            al=M.CENTRO)
    ancho_desc = sum(ws.column_dimensions[L(x)].width or 9 for x in range(2, c["cant"]))
    l_desc = math.ceil(len(it["descripcion"] or "") * 1.1 / max(10, ancho_desc - 2))
    ws.row_dimensions[f].height = max(55.2, l_desc * 15 + 8)
    return f, c, dat


def _final(ws, W, f_ult, fila_titulos):
    """Logos e impresion, como la memoria general."""
    M._logo(ws, "logo_izquierda.png", 0, 10, 16, 288, 77)
    x = sum(M._ancho_px(ws, c) for c in range(1, W + 1))
    x_cons = x - 119 - 8
    x_sup = x_cons - 205 - 18
    for nombre, xx, yy, an, al_ in (("logo_supervision.png", x_sup, 13, 205, 56),
                                    ("logo_contratista.png", x_cons, 5, 119, 73)):
        col, acum_ = 0, 0
        while col < W and acum_ + M._ancho_px(ws, col + 1) <= xx:
            acum_ += M._ancho_px(ws, col + 1)
            col += 1
        M._logo(ws, nombre, col, max(0, xx - acum_), yy, an, al_)
    ws.print_area = f"A1:{L(W)}{f_ult}"
    if fila_titulos:
        ws.print_title_rows = f"{fila_titulos[0]}:{fila_titulos[1]}"
    ws.page_setup.orientation = "portrait" if W <= 14 else "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left, ws.page_margins.right = 0.3937, 0.1181
    ws.page_margins.top, ws.page_margins.bottom = 0.7874, 0.1969
    ws.page_margins.header = ws.page_margins.footer = 0.315
    ws.sheet_view.view = "pageBreakPreview"


def _anchos(ws, W, anchos, codigos=()):
    if codigos:
        anchos[1] = max(anchos.get(1, 13.5), max(len(str(c or "")) for c in codigos) * 1.25 + 3)
    for c in range(1, W + 1):
        ws.column_dimensions[L(c)].width = round(anchos.get(c, 13.5), 1)


# ------------------------------------------------------------------ formato de tabla
def hoja_tabla(wb, p, it, desde, hasta, enc, pl, solo_contractual=False, usados=None, estimacion=False):
    cortes = [(iso_a_fecha(desde), iso_a_fecha(hasta))] if estimacion else M.semanas(desde, hasta, p.inicio_semana)
    d = Datos(p, it, desde, hasta, cortes, solo_contractual)
    usados = set() if usados is None else usados
    ws = wb.create_sheet(M.nombre_hoja(it, usados))
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 70
    uni = it["unidad"] or ""
    nom = dict(NOMBRES)
    nom.update(pl["nombres"])
    nom = {k: v.replace("{u}", uni) for k, v in nom.items()}
    etqs = [(t.replace("{u}", uni), tipo) for t, tipo in pl["etq"]]
    modo_nv = pl["etq_nivel"]

    # ---- filas: bloques [(titulo, [fila])]; fila = {"etq": {tipo: texto}, "f": filtro, "grupo": clave para combinar}
    filas_modo = pl["filas"]
    tipos_bloque = d.tipos if (d.con_tipos and filas_modo != "seccion" and not pl.get("tipos_en_filas")) else [None]
    columnas_nv = []          # etiquetas de nivel para las columnas divididas
    split = set(pl["split"])
    if split:
        nvs = d.niveles_usados()
        etq_de_nv = {nv: _etq_nivel(d.ubic[nv]["nv_nombre"], modo_nv) for nv in nvs}
        # un area exterior sin niveles va en la primera columna (como la memoria general)
        con_nombre = [e for nv, e in etq_de_nv.items() if d.ubic[nv]["nv_nombre"]]
        orden = sorted(set(con_nombre), key=lambda e: M._orden_nivel(next((d.ubic[nv]["nv_nombre"] for nv in nvs
                                                                            if etq_de_nv[nv] == e), e)))
        columnas_nv = orden or ["ÚNICA"]
        etq_de_nv = {nv: (e if d.ubic[nv]["nv_nombre"] else columnas_nv[0]) for nv, e in etq_de_nv.items()}
    else:
        etq_de_nv = {}

    def nvs_de_col(ed, col):
        return {nv for nv, e in etq_de_nv.items() if e == col and d.ubic[nv]["edificio"] == ed}

    bloques = []
    for tp in tipos_bloque:
        ft = {"tp": tp} if tp is not None else {}
        filas = []
        if filas_modo == "edificio":
            if pl.get("tipos_en_filas") and d.con_tipos:
                for ed in d.edificios_usados():
                    for t2 in d.tipos:
                        if d.hay(ed=ed, tp=t2):
                            filas.append({"etq": {"ed": d.nombre_ed[ed], "tipo": t2}, "f": dict(ed=ed, tp=t2),
                                          "ed": ed})
            else:
                for ed in d.edificios_usados(**ft):
                    filas.append({"etq": {"ed": d.nombre_ed[ed]}, "f": dict(ed=ed, **ft), "ed": ed})
        elif filas_modo == "ubicacion_plana":
            for nv in d.niveles_usados(**ft):
                filas.append({"etq": {"ubic": d.ubic[nv]["etiqueta"], "tipo": tp or ""}, "f": dict(nv=nv, **ft),
                              "ed": d.ubic[nv]["edificio"]})
        elif filas_modo == "ubicacion":
            for nv in d.niveles_usados(**ft):
                u = d.ubic[nv]
                filas.append({"etq": {"ed": d.nombre_ed[u["edificio"]], "nivel_num": _num_nivel(u["nv_nombre"])},
                              "f": dict(nv=nv, **ft), "grupo": u["edificio"], "ed": u["edificio"]})
        elif filas_modo == "tramo":
            for nv in d.niveles_usados(**ft):
                subs = sorted({k[2] for k in d.claves if k[1] == nv and d._ok(k, ft) and d.hay(nv=nv, tr=k[2], **ft)},
                              key=lambda t: (t != "", d.orden_tr.get(t, 0)))
                u = d.ubic[nv]
                for tr in subs:
                    filas.append({"etq": {"ubic": u["etiqueta"], "tramo": d.tramos[tr]["nombre"] if tr else "—",
                                          "nivel_rom": M.etiqueta_nivel(u["nv_nombre"]) if u["nv_nombre"] else "—"},
                                  "f": dict(nv=nv, tr=tr, **ft), "grupo": nv, "ed": u["edificio"]})
        elif filas_modo == "seccion":
            # secciones por nivel; dentro, cada edificio con sus tipos / equipos / tramos
            secciones = {}
            for nv in d.niveles_usados():
                secciones.setdefault(_etq_nivel(d.ubic[nv]["nv_nombre"], None) or d.ubic[nv]["etiqueta"], []).append(nv)
            for sec in sorted(secciones, key=M._orden_nivel):
                filas_s = []
                for nv in sorted(secciones[sec], key=lambda x: d.orden_nv.get(x, 0)):
                    u = d.ubic[nv]
                    if d.con_tipos:
                        subs = [("tp", t2, t2) for t2 in d.tipos if d.hay(nv=nv, tp=t2)]
                    elif d.equipos:
                        subs = [("eq", e, d.equipos[e]["nombre"]) for e in sorted(
                            {k[3] for k in d.claves if k[1] == nv and d.hay(nv=nv, eq=k[3])},
                            key=lambda e: (e == "", d.orden_eq.get(e, 0)))]
                    elif d.tramos:
                        subs = [("tr", t, d.tramos[t]["nombre"] if t else "—") for t in sorted(
                            {k[2] for k in d.claves if k[1] == nv and d.hay(nv=nv, tr=k[2])},
                            key=lambda t: (t == "", d.orden_tr.get(t, 0)))]
                    else:
                        subs = [(None, None, "—")]
                    for campo, val, texto in subs:
                        f = dict(nv=nv)
                        if campo:
                            f[campo] = val
                        filas_s.append({"etq": {"ed": u["ed_nombre"], "sub": texto if texto else "—"}, "f": f,
                                        "grupo": nv, "ed": u["edificio"], "nv_grupo": nv})
                bloques.append((sec, filas_s))
            break
        if filas_modo != "seccion":
            bloques.append((tp if (tp is not None and d.con_tipos) else None, filas))

    # ---- columnas
    hitos_cols = []
    if pl["hitos_per"] and d.hitos and not d.con_tipos:
        hitos_cols = [h for h in d.hitos]
    cols = []      # (clave, subtitulo, extra)
    for k, (t, tipo) in enumerate(etqs):
        cols.append(("etq", tipo, None))
    n_sem = len(cortes)

    def grupo(clave):
        if clave in split:
            res = [(clave, c_, c_) for c_ in columnas_nv]
            if clave == "dis" and pl["total_dis"] and len(columnas_nv) > 1:
                res.append((clave, "TOTAL", "TOTAL"))
            return res
        return [(clave, None, None)]

    cols += grupo("dis")
    if "pu" in pl["extras"]:
        cols.append(("pu", None, None))
    if "costo" in pl["extras"]:
        cols.append(("costo", None, None))
    if "monto_per" in pl["extras"]:
        cols.append(("monto_per", None, None))
    cols += grupo("ant")
    for k in range(n_sem):
        if hitos_cols and "per" not in split:
            for h in hitos_cols:
                cols.append((("per", k), f"EJECUTADO {h['nombre']}".upper(), ("hito", h["id"])))
            cols.append((("per", k), "TOTAL EJECUTADO", "TOTAL_H"))
        else:
            for _c, sub, nvc in grupo("per"):
                cols.append((("per", k), sub, nvc))
            if "per" in split and len(columnas_nv) > 1:
                cols.append((("per", k), "TOTAL", "TOTAL"))
    cols += grupo("acu")
    cols += grupo("pen")
    if "pct" in pl["extras"]:
        if "pct" in split:
            cols += grupo("pct")
            if len(columnas_nv) > 1:
                cols.append(("pct", "TOTAL", "TOTAL"))
        else:
            cols.append(("pct", None, None))
    if "monto_pen" in pl["extras"]:
        cols.append(("monto_pen", None, None))
    # acumulado / pendiente divididos por nivel: tambien su TOTAL
    for clave in ("acu", "pen"):
        if clave in split and len(columnas_nv) > 1:
            pos = max(j for j, c_ in enumerate(cols, 1) if c_[0] == clave)
            cols.insert(pos, (clave, "TOTAL", "TOTAL"))
    W = max(len(cols), 12)
    ncol = len(cols)

    # anchos
    anchos = {}
    for j, (cl, sub, ex) in enumerate(cols, start=1):
        if cl == "etq":
            textos = [str(f_["etq"].get(sub, "")) for _t, fs in bloques for f_ in fs] + [etqs[j - 1][0]]
            if j == 1 and len(etqs) == 1:          # «SUB TOTAL …» va en la primera columna
                textos += [f"SUB TOTAL {t_}" for t_, _fs in bloques if t_]
            anchos[j] = max(8.0, min(52.0, max(len(x) for x in textos) * 1.35 + 4))
        else:
            anchos[j] = 15.0 if (sub or "").startswith("EJECUTADO") else 14.0
    _anchos(ws, W, anchos, [it["codigo"]] + [a.get("codigo") for a in _ancestros(p, it)])
    f_item, ci, dat = _encabezado(ws, p, it, W, desde, hasta, enc)

    # ---- titulo de la tabla y encabezados
    f = f_item + 3
    th = dict(font=M._f(10, True, blanca=True, nombre="Calibri"), fill=M._r(M.AZUL), al=M.CENTRO, borde=M._caja())
    if pl["titulo"]:
        ws.row_dimensions[f].height = 24
        M._unir(ws, f, 1, f, ncol, pl["titulo"], font=M._f(14, True, blanca=True, nombre="Calibri"),
                fill=M._r(M.AZUL_RES), borde=M._caja(), al=M.CENTRO)
        f += 1
    dos = any(sub is not None for cl, sub, ex in cols if cl != "etq")
    f_h1 = f
    titulos = {"dis": nom["dis"], "ant": nom["ant"], "acu": nom["acu"], "pen": nom["pen"],
               "pct": nom.get("pct", "% AVANCE"), "pu": "COSTO UNITARIO", "costo": "COSTO TOTAL",
               "monto_per": "MONTO EJECUTADO ESTE PERÍODO", "monto_pen": "MONTO PENDIENTE"}

    def titulo_col(cl):
        if isinstance(cl, tuple):
            k = cl[1]
            if estimacion:
                return nom["per"]
            d0, d1 = cortes[k]
            return f"{nom['per']}\nSEMANA {M.texto_rango(d0, d1)}"
        return titulos.get(cl, "")

    ws.row_dimensions[f].height = 42
    if dos:
        ws.row_dimensions[f + 1].height = 30
    j = 1
    while j <= ncol:
        cl, sub, ex = cols[j - 1]
        if cl == "etq":
            M._unir(ws, f, j, f + (1 if dos else 0), j, etqs[j - 1][0], **th)
            j += 1
            continue
        k = j
        while k < ncol and cols[k][0] == cl:
            k += 1
        # columnas j..k con el mismo grupo
        if dos and any(cols[x - 1][1] is not None for x in range(j, k + 1)):
            M._unir(ws, f, j, f, k, titulo_col(cl), **th)
            for x in range(j, k + 1):
                M._unir(ws, f + 1, x, f + 1, x, cols[x - 1][1] or "", **th)
        else:
            M._unir(ws, f, j, f + (1 if dos else 0), k, titulo_col(cl), **th)
        j = k + 1
    f += 2 if dos else 1
    f_h2 = f - 1

    # ---- cuerpo
    col_de = {}
    for j, c_ in enumerate(cols, start=1):
        col_de.setdefault(c_[0], []).append(j)
    dis_cols = [j for j, c_ in enumerate(cols, 1) if c_[0] == "dis" and c_[2] != "TOTAL"]
    dis_tot = next((j for j, c_ in enumerate(cols, 1) if c_[0] == "dis" and c_[2] == "TOTAL"), None)
    ant_cols = col_de.get("ant", [])

    def dis_ref(f_, nvc=None):
        if nvc is not None and nvc != "TOTAL":
            jj = next((j for j in dis_cols if cols[j - 1][2] == nvc), None)
            return f"{L(jj)}{f_}" if jj else "0"
        if dis_tot:
            return f"{L(dis_tot)}{f_}"
        return "+".join(f"{L(j)}{f_}" for j in dis_cols) or "0"

    def ejec_refs(f_, nvc):
        """referencias a anterior + periodos del nivel nvc (o del total)"""
        refs = []
        # anterior
        if "ant" in split and nvc not in (None, "TOTAL"):
            refs += [f"{L(j)}{f_}" for j in ant_cols if cols[j - 1][2] == nvc]
        else:
            refs += [f"{L(j)}{f_}" for j in ant_cols if cols[j - 1][2] not in ("TOTAL",)] if not (
                "ant" in split and nvc in (None, "TOTAL")) else [f"{L(j)}{f_}" for j in ant_cols]
        # periodos
        for k in range(n_sem):
            js = [j for j, c_ in enumerate(cols, 1) if c_[0] == ("per", k)]
            if hitos_cols and "per" not in split:
                refs += [f"{L(j)}{f_}" for j in js if cols[j - 1][2] == "TOTAL_H"]
            elif "per" in split and nvc not in (None, "TOTAL"):
                refs += [f"{L(j)}{f_}" for j in js if cols[j - 1][2] == nvc]
            else:
                refs += [f"{L(j)}{f_}" for j in js if cols[j - 1][2] != "TOTAL"]
        return refs

    blanco = M._r(M.BLANCO)
    num_al = Alignment(horizontal="right", vertical="center")
    cen = Alignment(horizontal="center", vertical="center", wrap_text=True)
    subtotales = []
    contador = 0

    def escribir_fila(f_, fila, negrita=False, relleno=None, formula_rango=None):
        """fila de datos (formula_rango=None) o de subtotal/total (formula_rango = lista de filas a sumar)."""
        fuente = M._f(12, negrita)
        rel = relleno or blanco
        for j, (cl, sub, ex) in enumerate(cols, start=1):
            cel_borde = M._caja()
            if cl == "etq":
                if formula_rango is None:
                    v = fila["etq"].get(sub, "")
                    if sub == "num":
                        v = fila["etq"].get("num", "")
                    if sub == "unidad":
                        v = uni
                    if sub == "dis_ed":
                        v = None
                    M._celda(ws, f_, j, v, fuente, rel, cel_borde, "@" if sub != "dis_ed" else M.NUM,
                             cen if sub != "dis_ed" else num_al)
                else:
                    M._celda(ws, f_, j, None, fuente, rel, cel_borde, None, cen)
                continue
            if formula_rango is not None:
                if cl == "pct":
                    nvc = ex if ex not in (None, "TOTAL") else None
                    acu_j = _col_acu(nvc)
                    v = f"=IFERROR({L(acu_j)}{f_}/{dis_ref(f_, nvc)},0)" if acu_j else None
                elif cl == "pu":
                    v = None
                elif formula_rango:
                    contiguo = formula_rango == list(range(formula_rango[0], formula_rango[-1] + 1))
                    v = (f"=SUM({L(j)}{formula_rango[0]}:{L(j)}{formula_rango[-1]})" if contiguo else
                         "=" + "+".join(f"{L(j)}{x}" for x in formula_rango))
                else:
                    v = 0
                M._celda(ws, f_, j, v, fuente, rel, cel_borde, PCT if cl == "pct" else M.NUM, num_al)
                continue
            filtro = dict(fila["f"])
            nvc = ex if (ex not in (None, "TOTAL", "TOTAL_H") and not isinstance(ex, tuple)) else None
            if nvc is not None:
                filtro["nvs"] = nvs_de_col(fila["ed"], nvc)
            v = None
            fmt = M.NUM
            if cl == "dis":
                if ex == "TOTAL":
                    v = "=ROUND(" + "+".join(f"{L(x)}{f_}" for x in dis_cols) + ",2)"
                else:
                    x = d.diseno(**filtro) if (nvc is None or filtro["nvs"]) else None
                    v = round(x, 4) if x is not None else (0 if d.con_diseno else None)
            elif cl == "ant" or isinstance(cl, tuple):
                if ex in ("TOTAL", "TOTAL_H"):
                    js = [x for x, c_ in enumerate(cols, 1) if c_[0] == cl and c_[2] not in ("TOTAL", "TOTAL_H")]
                    v = "=" + "+".join(f"{L(x)}{f_}" for x in js)
                else:
                    k_ = 0 if cl == "ant" else cl[1] + 1
                    hito = ex[1] if isinstance(ex, tuple) else None
                    if nvc is not None and not filtro["nvs"]:
                        x = 0.0
                    else:
                        x = d.serie(hito=hito, **filtro)[k_]
                    v = round(x, 4) if abs(x) > 1e-9 else 0
            elif cl == "acu":
                refs = ejec_refs(f_, nvc)
                v = "=" + "+".join(refs) if refs else 0
            elif cl == "pen":
                acu_j = _col_acu(nvc)
                v = f"={dis_ref(f_, nvc)}-{L(acu_j)}{f_}" if (acu_j and d.con_diseno) else None
            elif cl == "pct":
                acu_j = _col_acu(nvc)
                v = f"=IFERROR({L(acu_j)}{f_}/{dis_ref(f_, nvc)},0)" if (acu_j and d.con_diseno) else None
                fmt = PCT
            elif cl == "pu":
                v = it.get("pu") or 0
            elif cl == "costo":
                v = f"={dis_ref(f_)}*{L(col_de['pu'][0])}{f_}" if "pu" in col_de else None
            elif cl == "monto_per":
                per_tot = [x for x, c_ in enumerate(cols, 1) if isinstance(c_[0], tuple)
                           and c_[2] in (None, "TOTAL", "TOTAL_H")]
                if hitos_cols or "per" in split:
                    per_tot = [x for x in per_tot if cols[x - 1][2] in ("TOTAL", "TOTAL_H")]
                v = ("=(" + "+".join(f"{L(x)}{f_}" for x in per_tot) + f")*{L(col_de['pu'][0])}{f_}"
                     if per_tot and "pu" in col_de else None)
            elif cl == "monto_pen":
                pen_j = next((x for x, c_ in enumerate(cols, 1) if c_[0] == "pen" and c_[2] in (None, "TOTAL")), None)
                v = f"={L(pen_j)}{f_}*{L(col_de['pu'][0])}{f_}" if (pen_j and "pu" in col_de) else None
            M._celda(ws, f_, j, v, fuente, rel, cel_borde, fmt, num_al)

    def _col_acu(nvc):
        for j, c_ in enumerate(cols, 1):
            if c_[0] == "acu" and (c_[2] == nvc or (nvc is None and c_[2] in (None, "TOTAL"))):
                return j
        if nvc is None:
            js = [j for j, c_ in enumerate(cols, 1) if c_[0] == "acu"]
            return js[0] if len(js) == 1 else None
        return None


    filas_datos_todas = []
    for tit, filas in bloques:
        if tit is not None:
            ws.row_dimensions[f].height = 24
            M._unir(ws, f, 1, f, ncol, tit, font=M._f(12, True, nombre="Calibri"), fill=M._r(M.SALMON),
                    borde=M._caja(), al=M.CENTRO)
            f += 1
        f_ini = f
        grupos_f = {}
        for fila in filas:
            contador += 1
            fila["etq"]["num"] = contador
            ws.row_dimensions[f].height = 18
            escribir_fila(f, fila)
            if fila.get("grupo") is not None:
                grupos_f.setdefault(fila["grupo"], []).append(f)
            f += 1
        # celdas combinadas en vertical (edificio / sector) y diseño del bloque (HVAC)
        for g, fs in grupos_f.items():
            for j, (cl, sub, ex) in enumerate(cols, 1):
                if cl == "etq" and sub in ("ed", "ubic") and len(fs) > 1 and fs == list(range(fs[0], fs[-1] + 1)):
                    v = ws.cell(fs[0], j).value
                    M._unir(ws, fs[0], j, fs[-1], j, v, font=M._f(12, True), fill=blanco, borde=M._caja(), al=cen)
                if cl == "etq" and sub == "dis_ed":
                    col_d = dis_tot or (dis_cols[0] if dis_cols else None)
                    v = f"=SUM({L(col_d)}{fs[0]}:{L(col_d)}{fs[-1]})" if (col_d and d.con_diseno) else None
                    M._unir(ws, fs[0], j, fs[-1], j, v, font=M._f(12), fill=blanco, borde=M._caja(), al=num_al,
                            fmt=M.NUM)
        f_fin = f - 1
        filas_datos_todas.append((f_ini, f_fin))
        if len(bloques) > 1 or filas_modo == "seccion":
            ws.row_dimensions[f].height = 19.2
            escribir_fila(f, None, True, M._r(M.CELESTE), list(range(f_ini, f_fin + 1)))
            M._unir(ws, f, 1, f, max(1, len(etqs)), f"SUB TOTAL {tit}" if tit else "SUB TOTAL",
                    font=M._f(12, True), fill=M._r(M.CELESTE), borde=M._caja(), al=cen)
            subtotales.append(f)
            f += 1
    # total
    f_total = f
    ws.row_dimensions[f].height = 26
    rango = subtotales if subtotales else [x for a, b in filas_datos_todas for x in range(a, b + 1)]
    escribir_fila(f, None, True, M._r(M.CELESTE), rango)
    M._unir(ws, f, 1, f, max(1, len(etqs)), "TOTAL", font=M._f(13, True), fill=M._r(M.CELESTE), borde=M._caja(),
            al=cen)
    # sin datos fijos: lo pendiente sale de la cantidad total de la partida
    if not d.con_diseno:
        for j, c_ in enumerate(cols, 1):
            if c_[0] == "dis" and c_[2] in (None, "TOTAL"):
                ws.cell(f_total, j).value = f"=ROUND({L(ci['cant'])}{f_item},2)"
            if c_[0] == "pen" and c_[2] in (None, "TOTAL"):
                acu_j = _col_acu(None)
                ws.cell(f_total, j).value = f"={dis_ref(f_total)}-{L(acu_j)}{f_total}" if acu_j else None

    # ---- fila del item: anterior, periodo y acumulado de la partida (formulas)
    ant_ref = "+".join(f"{L(j)}{f_total}" for j in ant_cols if cols[j - 1][2] != "TOTAL") or "0"
    per_refs = []
    for k in range(n_sem):
        js = [j for j, c_ in enumerate(cols, 1) if c_[0] == ("per", k)]
        tot = [j for j in js if cols[j - 1][2] in ("TOTAL", "TOTAL_H")]
        per_refs += [f"{L(j)}{f_total}" for j in (tot or js)]
    M._unir(ws, f_item, ci["aa"], f_item, ci["aa"], f"=ROUND({ant_ref},2)", fmt=M.NUM, **dat)
    M._unir(ws, f_item, ci["ep"], f_item, ci["ep"], "=ROUND(" + ("+".join(per_refs) or "0") + ",2)", fmt=M.NUM, **dat)
    M._unir(ws, f_item, ci["ea"], f_item, ci["ea"], f"={L(ci['aa'])}{f_item}+{L(ci['ep'])}{f_item}", fmt=M.NUM, **dat)
    _final(ws, W, f_total, (f_h1, f_h2))
    return ws


# ------------------------------------------------------------------ formato de pago por hitos
def hoja_hitos(wb, p, it, desde, hasta, enc, solo_contractual=False, usados=None, estimacion=False):
    cortes = [(iso_a_fecha(desde), iso_a_fecha(hasta))] if estimacion else M.semanas(desde, hasta, p.inicio_semana)
    d = Datos(p, it, desde, hasta, cortes, solo_contractual)
    usados = set() if usados is None else usados
    ws = wb.create_sheet(M.nombre_hoja(it, usados))
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 70
    n_sem = len(cortes)
    # columnas: A N° | B-C HITO | D PORCENTAJE | E ANTERIOR | semanas | ACUMULADO | PENDIENTE
    c_pct, c_ant = 4, 5
    c_sem = [6 + k for k in range(n_sem)]
    c_acu = 6 + n_sem
    c_pen = c_acu + 1
    W = max(c_pen + 1, 12)
    anchos = {1: 9.0, 2: 17.4, 3: 24.7, 4: 17.7, 5: 17.0}
    for c in [*c_sem, c_acu, c_pen]:
        anchos[c] = 16.0
    _anchos(ws, W, anchos, [it["codigo"]] + [a.get("codigo") for a in _ancestros(p, it)])
    f_item, ci, dat = _encabezado(ws, p, it, W, desde, hasta, enc)
    cont = it["cantidad"] or 0.0
    total_cont = d.diseno() if d.con_diseno else None
    base = total_cont if total_cont else cont

    f = f_item + 2
    th = dict(font=M._f(10, True, nombre="Calibri"), fill=M._r(M.SALMON_B), borde=M._caja(), al=M.CENTRO)
    ws.row_dimensions[f].height = 22
    M._unir(ws, f, 1, f, c_pen, "CONTROL DE PAGO POR HITOS", font=M._f(11, True, blanca=True, nombre="Calibri"),
            fill=M._r(M.AZUL), borde=M._caja(), al=M.CENTRO)
    f += 1
    f_h = f
    ws.row_dimensions[f].height = 30
    M._unir(ws, f, 1, f, 1, "N°", **th)
    M._unir(ws, f, 2, f, 3, "HITO", **th)
    M._unir(ws, f, c_pct, f, c_pct, "PORCENTAJE", **th)
    M._unir(ws, f, c_ant, f, c_ant, "ANTERIOR", **th)
    for k, c in enumerate(c_sem):
        M._unir(ws, f, c, f, c, "ESTE PERÍODO" if estimacion else "SEMANA\n" + M.texto_rango(*cortes[k]), **th)
    M._unir(ws, f, c_acu, f, c_acu, "ACUMULADO", **th)
    M._unir(ws, f, c_pen, f, c_pen, "PENDIENTE", **th)
    f += 1
    f_h0 = f
    dato = dict(font=M._f(10, nombre="Calibri"), borde=M._caja(), al=M.CENTRO)
    hitos = d.hitos or HITOS_ESTANDAR
    for k, h in enumerate(hitos, start=1):
        ws.row_dimensions[f].height = 22
        M._unir(ws, f, 1, f, 1, k, font=M._f(10, True, nombre="Calibri"), borde=M._caja(), al=M.CENTRO)
        M._unir(ws, f, 2, f, 3, h["nombre"], font=M._f(10, True, nombre="Calibri"), borde=M._caja(), al=M.IZQ)
        M._unir(ws, f, c_pct, f, c_pct, round(h["pct"] / 100.0, 6), fmt=PCT, **dato)
        # avance del hito como fraccion de la partida: lo ponderado del hito / cantidad de la partida
        serie = d.serie(hito=h["id"])
        for c, x in zip([c_ant] + c_sem, serie):
            M._unir(ws, f, c, f, c, round(x / base, 6) if base else 0, fmt=PCT, **dato)
        M._unir(ws, f, c_acu, f, c_acu, "=" + "+".join(f"{L(c)}{f}" for c in [c_ant] + c_sem), fmt=PCT, **dato)
        M._unir(ws, f, c_pen, f, c_pen, f"={L(c_pct)}{f}-{L(c_acu)}{f}", fmt=PCT, **dato)
        f += 1
    f_h1 = f - 1
    tot = dict(font=M._f(10, True, nombre="Calibri"), fill=M._r(DDEBF7), borde=M._caja(), al=M.CENTRO)
    M._unir(ws, f, 1, f, 3, "TOTAL", **tot)
    for c in [c_pct, c_ant, *c_sem, c_acu, c_pen]:
        M._unir(ws, f, c, f, c, f"=SUM({L(c)}{f_h0}:{L(c)}{f_h1})" if f_h1 >= f_h0 else 0, fmt=PCT, **tot)
    f_tot = f
    f += 1
    # unidades base del hito 1 (cuando el despacho es parcial)
    h1 = hitos[0] if hitos else None
    u1 = d.unid_hito.get(h1["id"], [0.0] * d.n) if h1 else [0.0] * d.n
    u1_acu = sum(u1)
    u1_per = sum(u1[1:])
    parcial = h1 is not None and base and 1e-9 < u1_acu < base - 1e-9
    if parcial:
        M._unir(ws, f, 1, f, 3, f"{it['unidad'] or 'UND'} despachadas este período (base Hito 1):",
                font=M._f(10, True, nombre="Calibri"), al=M.IZQ)
        M._celda(ws, f, c_pct, round(u1_per, 4), M._f(10, True, nombre="Calibri"), M._r(FFF2CC), M._caja(), M.NUM,
                 M.CENTRO)
        M._celda(ws, f, c_ant, it["unidad"] or "UND", M._f(10, nombre="Calibri"), None, None, None, M.CENTRO)
        f += 1
    # cantidad de la partida
    ws.row_dimensions[f].height = 30
    M._unir(ws, f, 1, f, 3, "DESCRIPCIÓN", **th)
    M._unir(ws, f, c_pct, f, c_pct, "CANTIDAD CONTRACTUAL", **th)
    M._unir(ws, f, c_ant, f, c_ant, "ANTERIOR", **th)
    for k, c in enumerate(c_sem):
        M._unir(ws, f, c, f, c, "ESTE PERÍODO" if estimacion else "SEMANA\n" + M.texto_rango(*cortes[k]), **th)
    M._unir(ws, f, c_acu, f, c_acu, "ACUMULADO", **th)
    M._unir(ws, f, c_pen, f, c_pen, "PENDIENTE", **th)
    f += 1
    f_c = f
    import math
    ws.row_dimensions[f].height = max(30, math.ceil(len(it["descripcion"] or "") / 60) * 13 + 6)
    M._unir(ws, f, 1, f, 3, it["descripcion"], font=M._f(9, nombre="Calibri"), fill=M._r(DDEBF7), borde=M._caja(),
            al=M.IZQ)
    cd = dict(font=M._f(10, True, nombre="Calibri"), fill=M._r(DDEBF7), borde=M._caja(), al=M.CENTRO)
    M._unir(ws, f, c_pct, f, c_pct, round(base, 4), fmt=M.NUM, **cd)
    for c in [c_ant, *c_sem]:
        M._unir(ws, f, c, f, c, f"=ROUND({L(c_pct)}{f}*{L(c)}{f_tot},4)", fmt=M.NUM, **cd)
    M._unir(ws, f, c_acu, f, c_acu, "=" + "+".join(f"{L(c)}{f}" for c in [c_ant, *c_sem]), fmt=M.NUM, **cd)
    M._unir(ws, f, c_pen, f, c_pen, f"={L(c_pct)}{f}-{L(c_acu)}{f}", fmt=M.NUM, **cd)
    f += 3
    # requisitos y nota
    M._unir(ws, f, 1, f, c_pen, "REQUISITOS POR HITO", font=M._f(11, True, blanca=True, nombre="Calibri"),
            fill=M._r(M.AZUL), borde=M._caja(), al=M.CENTRO)
    f += 1
    for k, h in enumerate(hitos, start=1):
        req = REQUISITOS.get(_norm(h["nombre"]))
        ws.row_dimensions[f].height = 32 if req and len(req) > 110 else 22
        M._unir(ws, f, 1, f, 3, f"HITO {k} · {h['nombre']}", font=M._f(10, True, nombre="Calibri"),
                fill=M._r(M.SALMON_B), borde=M._caja(), al=M.IZQ)
        M._unir(ws, f, 4, f, c_pen, req or "", font=M._f(10, nombre="Calibri"), borde=M._caja(), al=M.IZQ)
        f += 1
    f += 1
    if h1 is not None and _norm(h1["nombre"]).startswith("despacho"):
        pct1 = h1["pct"]
        pct_t = f"{pct1:g}"
        sobre = (f" sobre {u1_acu:g} de {base:g} {it['unidad'] or 'UND'} despachadas" if parcial else "")
        nota = (f"NOTA : Pago del {pct_t} % en concepto del Hito 1 según Enmienda No.1{sobre} (Equipos "
                "electromecánicos y subcontratos especializados, se adjunta la información correspondiente)")
        ws.row_dimensions[f].height = 15
        ws.row_dimensions[f + 1].height = 15
        M._unir(ws, f, 1, f + 1, c_pen, nota, font=M._f(10, nombre="Calibri"), fill=M._r(DDEBF7), borde=M._caja(),
                al=M.IZQ)
        f += 2
    # fila del item
    M._unir(ws, f_item, ci["aa"], f_item, ci["aa"], f"=ROUND({L(c_ant)}{f_c},2)", fmt=M.NUM, **dat)
    M._unir(ws, f_item, ci["ep"], f_item, ci["ep"], "=ROUND(" + "+".join(f"{L(c)}{f_c}" for c in c_sem) + ",2)",
            fmt=M.NUM, **dat)
    M._unir(ws, f_item, ci["ea"], f_item, ci["ea"], f"={L(ci['aa'])}{f_item}+{L(ci['ep'])}{f_item}", fmt=M.NUM, **dat)
    _final(ws, W, f - 1, (f_h, f_h))
    return ws


def hoja_propia(wb, p, it, desde, hasta, enc, solo_contractual=False, usados=None, estimacion=False):
    """La hoja con el formato propio de la actividad, o None si usa el formato general."""
    pl = plantilla_de(p, it)
    if pl is None:
        return None
    if pl[0] == "hitos":
        return hoja_hitos(wb, p, it, desde, hasta, enc, solo_contractual, usados, estimacion)
    return hoja_tabla(wb, p, it, desde, hasta, enc, pl[1], solo_contractual, usados, estimacion)

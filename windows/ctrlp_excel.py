"""
Ctrl-P - exportaciones a Excel
==============================
- Informe de avance (Estado al / Periodo): Resumen, Por partida, Detallado por
  ubicacion y Financiero. Columnas ocultables desde el programa y agrupadas con
  botones +/- dentro del Excel.
- Corte de estimacion: el presupuesto completo con su formato y orden, mas
  Cantidad de este periodo y Monto de este periodo (cobrable). Sin excedente.
"""

from copy import copy

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ctrlp_core import FILTROS, Calculo, ahora, buckets, completa, fecha_txt

AZUL = "4C6F9F"
GRIS_T = "EEF2F7"
GRIS_S = "F4F6F9"
ROJO = "B42318"
NUM = "#,##0.00"
PCT = "0.00%"
PCT_FIN = "0.000%"
F_TIT = Font(name="Calibri", bold=True, size=14, color="1E2730")
F_SUB = Font(name="Calibri", size=10, color="5F6B7A")
F_ENC = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
F_B = Font(name="Calibri", bold=True, size=10)
F_N = Font(name="Calibri", size=10)
F_ROJO = Font(name="Calibri", bold=True, size=10, color=ROJO)
VERDE = "067647"
F_VERDE = Font(name="Calibri", size=10, color=VERDE)
F_VERDE_B = Font(name="Calibri", bold=True, size=10, color=VERDE)
RELL_ENC = PatternFill("solid", fgColor=AZUL)
RELL_T0 = PatternFill("solid", fgColor="DCE3EC")
RELL_T = PatternFill("solid", fgColor=GRIS_T)
RELL_S = PatternFill("solid", fgColor=GRIS_S)
LINEA = Side(style="thin", color="D3DBE5")
BORDE = Border(bottom=LINEA)

COLUMNAS_OPCIONALES = [
    ("unidad", "Unidad"), ("cantidad", "Cantidad / contractual"), ("ant", "Acumulado anterior"),
    ("p_ant", "% anterior"), ("per", "Este período / del día"), ("p_per", "% del período"),
    ("tot", "Acumulado total"), ("p_tot", "% acumulado"), ("saldo", "Saldo"), ("obs", "Observaciones"),
]


def _encabezado(ws, p, titulo, sub, ncols):
    ws["A1"] = titulo
    ws["A1"].font = F_TIT
    ws["A2"] = p.nombre
    ws["A2"].font = Font(name="Calibri", bold=True, size=11)
    ws["A3"] = sub
    ws["A3"].font = F_SUB
    ws["A4"] = f"Presupuesto V{p.version_presupuesto} · generado {ahora()} por {p.usuario}"
    ws["A4"].font = F_SUB
    return 6


def _fila_enc(ws, fila, textos):
    for c, t in enumerate(textos, start=1):
        cel = ws.cell(row=fila, column=c, value=t)
        cel.font = F_ENC
        cel.fill = RELL_ENC
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 32


_ESTILOS = {}          # (libro, formato, letra, relleno) -> estilo ya armado (mucho mas rapido que armarlo celda por celda)


def _poner(ws, fila, col, valor, fmt=None, font=None, fill=None, rojo_neg=False):
    c = ws.cell(row=fila, column=col, value=valor)
    if rojo_neg and isinstance(valor, (int, float)) and valor < -1e-9:
        font = F_ROJO
    clave = (id(ws.parent), fmt or None, id(font or F_N), id(fill) if fill else 0)
    est = _ESTILOS.get(clave)
    if est is None:
        if fmt:
            c.number_format = fmt
        c.font = font or F_N
        if fill:
            c.fill = fill
        c.border = BORDE
        _ESTILOS[clave] = copy(c._style)
    else:
        c._style = copy(est)
    return c


def _contexto(p, modo, desde, hasta):
    if modo == "estado":
        return f"Estado al {fecha_txt(hasta)}"
    return f"Período del {fecha_txt(desde)} al {fecha_txt(hasta)}"


def _tiene_avance(r):
    return abs(r.get("tot", 0)) > 1e-9 or abs(r.get("per", 0)) > 1e-9


def _guardar(wb, ruta, guardar):
    if guardar:
        guardar(wb, ruta)
    else:
        wb.save(ruta)


def _avisador(progreso):
    """progreso(fraccion, texto) o nada. Devuelve aviso(desde, hasta, k, n, texto) que reparte la fraccion."""
    def aviso(f0, f1, k, n, texto):
        if progreso:                       # el dialogo ya limita cuantas veces se redibuja
            progreso(f0 + (f1 - f0) * (k + 1) / max(1, n), texto)
    return aviso


def exportar_informe(p, ruta, modo, desde, hasta, detalle="", visibles=None, filtro="completo",
                     solo_contractual=False, progreso=None, guardar=None):
    """modo: 'estado' (desde=hasta=fecha) o 'periodo'. detalle: '', 'dia', 'semana'.
    solo_contractual: en partidas repartidas por ubicacion, solo lo que cabe en la cantidad de cada ubicacion."""
    _ESTILOS.clear()
    aviso = _avisador(progreso)
    if progreso:
        progreso(0.02, "Calculando…")
    if modo == "estado":
        desde = hasta
    visibles = set(visibles or [k for k, _ in COLUMNAS_OPCIONALES])
    if filtro is True:
        filtro = "historial"
    elif not filtro:
        filtro = "completo"
    calc = Calculo(p, desde, hasta, solo_contractual)
    vis_ids = calc.visibles(filtro)
    cortes = buckets(desde, hasta, detalle) if (detalle and modo == "periodo") else []
    ex_buckets = []
    for k, (_, b0, b1) in enumerate(cortes):
        aviso(0.02, 0.05, k, len(cortes), "Calculando días y semanas…")
        ex_buckets.append(p.ejecutado_periodo(b0, b1, solo_contractual))
    obs = p.observaciones(desde, hasta) if "obs" in visibles else {}
    wb = openpyxl.Workbook()
    ctx = _contexto(p, modo, desde, hasta)
    if solo_contractual:
        ctx += " · solo lo contractual"
    if filtro != "completo":
        ctx += " · " + dict(FILTROS).get(filtro, "").lower()
    mon = p.get_meta("moneda", "US$") or "US$"

    # ------------------------------------------------------------ Resumen
    ws = wb.active
    ws.title = "Resumen"
    f = _encabezado(ws, p, "Informe de avance", ctx, 4)
    t = calc.totales()
    filas = [("Monto del presupuesto", t["total"], None),
             ("Avance financiero acumulado anterior", t["m_ant"], t["p_ant"]),
             ("Avance financiero " + ("del día" if modo == "estado" else "del período"), t["m_per"], t["p_per"]),
             ("Avance financiero acumulado total", t["m_tot"], t["p_tot"]),
             ("Saldo por ejecutar", t["total"] - t["m_tot"], t["p_saldo"]),
             ("Fuera de lo contractual (acumulado)" if solo_contractual else "Excedente no cobrable (acumulado)",
              t["m_exc"], None)]
    _fila_enc(ws, f, ["Concepto", f"Monto ({mon})", "% del presupuesto"])
    for k, (n, m, pc) in enumerate(filas, start=f + 1):
        _poner(ws, k, 1, n, font=F_B if k in (f + 3, f + 4) else None)
        _poner(ws, k, 2, round(m, 2), NUM)
        _poner(ws, k, 3, pc, PCT if pc is not None else None)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 18

    # ------------------------------------------------------------ Por partida
    ws = wb.create_sheet("Por partida")
    f = _encabezado(ws, p, "Avance por partida", ctx, 10)
    cols = [("codigo", "Código", 14), ("descripcion", "Descripción", 60)]
    if "unidad" in visibles:
        cols.append(("unidad", "Unidad", 9))
    if "cantidad" in visibles:
        cols.append(("cantidad", "Cantidad", 14))
    if modo == "periodo":
        if "ant" in visibles:
            cols.append(("ant", "Acum. anterior", 14))
        if "p_ant" in visibles:
            cols.append(("p_ant", "% anterior", 10))
    for k, (et, _, _) in enumerate(cortes):
        cols.append((f"b{k}", et, 12))
    if "per" in visibles or modo == "periodo":
        cols.append(("per", "Ejecutado del día" if modo == "estado" else "Sumatoria del período", 15))
    if "p_per" in visibles and modo == "periodo":
        cols.append(("p_per", "% período", 10))
    if "tot" in visibles:
        cols.append(("tot", "Acumulado a la fecha" if modo == "estado" else "Acum. total", 15))
    if "p_tot" in visibles:
        cols.append(("p_tot", "% acumulado", 11))
    if "saldo" in visibles:
        cols.append(("saldo", "Saldo", 14))
    if "obs" in visibles:
        cols.append(("obs", "Observaciones", 50))
    _fila_enc(ws, f, [c[1] for c in cols])
    for k, (_, _, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(k)].width = w
    fila = f + 1
    por_id = calc.por_id
    for n_it, it in enumerate(calc.items):
        aviso(0.05, 0.3, n_it, len(calc.items), "Hoja «Por partida»…")
        if it["id"] not in vis_ids:
            continue
        r = calc.res[it["id"]]
        if it["tipo"] == "T":
            fill = RELL_T0 if not it["padre"] else RELL_T
            for k, (key, _, _) in enumerate(cols, start=1):
                v = it["codigo"] if key == "codigo" else (it["descripcion"] if key == "descripcion" else None)
                _poner(ws, fila, k, v, font=F_B, fill=fill)
            ws.cell(row=fila, column=2).alignment = Alignment(indent=it["nivel"])
            fila += 1
            continue
        desc = it["descripcion"] + ("" if it["vigente"] else "  [fuera del presupuesto vigente]")
        fv = F_VERDE if completa(r) else None
        for k, (key, _, _) in enumerate(cols, start=1):
            if key == "codigo":
                _poner(ws, fila, k, it["codigo"], font=fv)
            elif key == "descripcion":
                c = _poner(ws, fila, k, desc, font=fv)
                c.alignment = Alignment(indent=it["nivel"], wrap_text=False)
            elif key == "unidad":
                _poner(ws, fila, k, it["unidad"], font=fv)
            elif key.startswith("b"):
                b = int(key[1:])
                _poner(ws, fila, k, round(sum(ex_buckets[b].get(it["id"], {}).values()), 4), NUM, font=fv)
            elif key == "obs":
                tx = []
                for (iid, nv), l in obs.items():
                    if iid == it["id"]:
                        tx += [f"{p.etiqueta_ubicacion(nv)} · {x}" for x in l]
                _poner(ws, fila, k, " | ".join(tx) or None, font=fv)
            elif key.startswith("p_"):
                _poner(ws, fila, k, r[key], PCT if r[key] is not None else None, font=fv)
            else:
                _poner(ws, fila, k, round(r[key], 4), NUM, font=fv, rojo_neg=(key == "saldo" and not fv))
        fila += 1
    _agrupar(ws, cols)
    # fijas: codigo .. % anterior (lo que sigue se desplaza)
    ws.freeze_panes = ws.cell(row=f + 1, column=1 + sum(1 for c in cols if c[0] in
                                                       ("codigo", "descripcion", "unidad", "cantidad", "ant", "p_ant")))

    # ------------------------------------------------------------ Detallado por ubicacion
    ws = wb.create_sheet("Detallado por ubicación")
    f = _encabezado(ws, p, "Avance detallado por ubicación", ctx, 12)
    dcols = [("codigo", "Código", 14), ("descripcion", "Actividad / ubicación", 56)]
    if "unidad" in visibles:
        dcols.append(("unidad", "Unidad", 9))
    if "cantidad" in visibles:
        dcols.append(("contractual", "Contractual", 14))
    if modo == "periodo" and "ant" in visibles:
        dcols.append(("ant", "Acum. anterior", 14))
    for k, (et, _, _) in enumerate(cortes):
        dcols.append((f"b{k}", et, 12))
    dcols.append(("per", "Ejecutado del día" if modo == "estado" else "Sumatoria del período", 15))
    if "tot" in visibles:
        dcols.append(("tot", "Acumulado a la fecha" if modo == "estado" else "Acum. total", 15))
    if "p_tot" in visibles:
        dcols += [("p_nivel", "% nivel", 10), ("p_edificio", "% edificio", 10), ("p_general", "% general", 10)]
    if "saldo" in visibles:
        dcols.append(("saldo", "Saldo", 14))
    _fila_enc(ws, f, [c[1] for c in dcols])
    for k, (_, _, w) in enumerate(dcols, start=1):
        ws.column_dimensions[get_column_letter(k)].width = w
    fila = f + 1

    def cub(iid, niveles):
        return {f"b{k}": sum(ex_buckets[k].get(iid, {}).get(nv, 0.0) for nv in niveles) for k in range(len(cortes))}
    for n_it, it in enumerate(calc.items):
        aviso(0.3, 0.6, n_it, len(calc.items), "Hoja «Detallado por ubicación»…")
        if it["id"] not in vis_ids:
            continue
        if it["tipo"] == "T":
            fill = RELL_T0 if not it["padre"] else RELL_T
            for k, (key, _, _) in enumerate(dcols, start=1):
                v = it["codigo"] if key == "codigo" else (it["descripcion"] if key == "descripcion" else None)
                _poner(ws, fila, k, v, font=F_B, fill=fill)
            ws.cell(row=fila, column=2).alignment = Alignment(indent=it["nivel"])
            fila += 1
            continue
        r = calc.res[it["id"]]
        vals = {"codigo": it["codigo"], "descripcion": it["descripcion"], "unidad": it["unidad"],
                "contractual": it["cantidad"], "ant": r["ant"], "per": r["per"], "tot": r["tot"],
                "p_general": r["p_tot"], "saldo": r["saldo"]}
        vals.update(cub(it["id"], set(r["nv_tot"]) | set(r["cont"])))
        _fila_detalle(ws, fila, dcols, vals, F_VERDE_B if completa(r) else F_B, None)
        ws.cell(row=fila, column=2).alignment = Alignment(indent=it["nivel"])
        fila += 1
        for g in calc.detalle_ubicaciones(it["id"]):
            for fl in g["filas"]:
                if not fl["tot"] and fl["contractual"] is None:
                    continue
                nombre = g["nombre"] if g["es_area"] else f"{g['nombre']} {fl['nombre']}"
                vv = {"descripcion": "    " + nombre, "unidad": it["unidad"], "contractual": fl["contractual"],
                      "ant": fl["ant"], "per": fl["per"], "tot": fl["tot"], "p_nivel": fl["p_nivel"],
                      "saldo": fl["saldo"]}
                vv.update(cub(it["id"], [fl["nivel"]]))
                listo = bool(fl["contractual"]) and fl["tot"] >= fl["contractual"] * (1 - 5e-5)
                _fila_detalle(ws, fila, dcols, vv, F_VERDE if listo else None, None)
                fila += 1
                for h in fl.get("hojas", []):
                    base = h["contractual"] if h["contractual"] else (fl["contractual"] if not h["tramo"] else None)
                    _fila_detalle(ws, fila, dcols, {
                        "descripcion": "        " + h["nombre"], "unidad": it["unidad"],
                        "contractual": f"{h['factor'] * 100:.2f} %" if h["hito"] else h["contractual"],
                        "ant": h["ant"], "per": h["per"], "tot": h["tot"],
                        "p_nivel": (h["tot"] / base) if base else None,
                        "saldo": (base - h["tot"]) if base else None}, None, None)
                    ws.row_dimensions[fila].outlineLevel = 1
                    fila += 1
            if not g["es_area"] and len(g["filas"]) > 0 and (g["tot"] or g["contractual"]):
                vv = {"descripcion": f"    Total {g['nombre']}", "unidad": it["unidad"], "contractual": g["contractual"],
                      "ant": g["ant"], "per": g["per"], "tot": g["tot"], "p_edificio": g["p_edificio"],
                      "p_general": g["p_general"], "saldo": g["saldo"]}
                vv.update(cub(it["id"], [x["nivel"] for x in g["filas"]]))
                _fila_detalle(ws, fila, dcols, vv, F_B, RELL_S)
                fila += 1
    _agrupar(ws, dcols)
    ws.freeze_panes = ws.cell(row=f + 1, column=1 + sum(1 for c in dcols if c[0] in
                                                       ("codigo", "descripcion", "unidad", "contractual", "ant")))

    # ------------------------------------------------------------ Financiero
    ws = wb.create_sheet("Financiero")
    f = _encabezado(ws, p, "Avance financiero", ctx, 11)
    fcols = [("codigo", "Código", 14), ("descripcion", "Descripción", 56), ("unidad", "Unidad", 9),
             ("pu", f"Precio unitario ({mon})", 14), ("total", f"Monto contrato ({mon})", 16),
             ("m_ant", f"Monto anterior ({mon})", 16),
             ("m_per", f"Monto {'del día' if modo == 'estado' else 'del período'} ({mon})", 16),
             ("m_tot", f"Monto acumulado ({mon})", 16), ("m_saldo", f"Saldo ({mon})", 16),
             ("a_per", "% aporte " + ("del día" if modo == "estado" else "del período"), 12),
             ("a_tot", "% aporte acumulado", 12), ("m_exc", (f"Fuera de lo contractual ({mon})" if solo_contractual else f"Excedente no cobrable ({mon})"), 16)]
    _fila_enc(ws, f, [c[1] for c in fcols])
    for k, (_, _, w) in enumerate(fcols, start=1):
        ws.column_dimensions[get_column_letter(k)].width = w
    fila = f + 1
    tp = calc.total_presu or 1
    for n_it, it in enumerate(calc.items):
        aviso(0.6, 0.68, n_it, len(calc.items), "Hoja «Financiero»…")
        r = calc.res[it["id"]]
        if it["id"] not in vis_ids:
            continue
        titulo = it["tipo"] == "T"
        fill = (RELL_T0 if not it["padre"] else RELL_T) if titulo else None
        fnt = F_B if titulo else (F_VERDE if completa(r) else None)
        vals = {"codigo": it["codigo"], "descripcion": it["descripcion"],
                "unidad": None if titulo else it["unidad"], "pu": None if titulo else it["pu"],
                "total": r["total"], "m_ant": r["m_ant"], "m_per": r["m_per"], "m_tot": r["m_tot"],
                "m_saldo": r["m_saldo"], "a_per": r["m_per"] / tp, "a_tot": r["m_tot"] / tp, "m_exc": r["m_exc"]}
        for k, (key, _, _) in enumerate(fcols, start=1):
            v = vals[key]
            fmt = PCT_FIN if key.startswith("a_") else (NUM if isinstance(v, float) else None)
            c = _poner(ws, fila, k, round(v, 6) if isinstance(v, float) else v, fmt, fnt, fill,
                       rojo_neg=(key == "m_saldo"))
            if key == "descripcion":
                c.alignment = Alignment(indent=it["nivel"])
        fila += 1
    t = calc.totales()
    tot = {"descripcion": "TOTAL DEL PROYECTO", "total": t["total"], "m_ant": t["m_ant"], "m_per": t["m_per"],
           "m_tot": t["m_tot"], "m_saldo": t["total"] - t["m_tot"], "a_per": t["p_per"], "a_tot": t["p_tot"],
           "m_exc": t["m_exc"]}
    for k, (key, _, _) in enumerate(fcols, start=1):
        v = tot.get(key)
        fmt = PCT_FIN if key.startswith("a_") else NUM
        _poner(ws, fila, k, round(v, 6) if isinstance(v, float) else v, fmt if v is not None else None,
               F_B, RELL_T0)
    ws.freeze_panes = ws.cell(row=f + 1, column=3)
    if progreso:
        progreso(0.7, "Guardando el archivo…")
    _guardar(wb, ruta, guardar)
    return ruta


def _fila_detalle(ws, fila, cols, vals, fnt, fill):
    for k, (key, _, _) in enumerate(cols, start=1):
        v = vals.get(key)
        if key.startswith("p_"):
            _poner(ws, fila, k, v, PCT if v is not None else None, fnt, fill)
        elif isinstance(v, (int, float)):
            _poner(ws, fila, k, round(v, 4), NUM, fnt, fill,
                   rojo_neg=(key == "saldo" and fnt not in (F_VERDE, F_VERDE_B)))
        else:
            _poner(ws, fila, k, v, None, fnt, fill)


def _agrupar(ws, cols):
    """Agrupa columnas de porcentajes y de detalle por dia/semana (botones +/- en Excel)."""
    keys = [c[0] for c in cols]
    for k, key in enumerate(keys, start=1):
        if key.startswith("p_") or (key.startswith("b") and key[1:].isdigit()):
            letra = get_column_letter(k)
            ws.column_dimensions.group(letra, letra, hidden=False, outline_level=1)
    ws.sheet_properties.outlinePr.summaryRight = False


def exportar_corte(p, ruta, desde, hasta, solo_contractual=False, progreso=None, guardar=None):
    """Corte de estimacion: presupuesto completo + cantidad y monto cobrables del periodo."""
    _ESTILOS.clear()
    calc = Calculo(p, desde, hasta, solo_contractual)
    mon = p.get_meta("moneda", "US$") or "US$"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Corte de estimación"
    ws["A1"] = p.nombre
    ws["A1"].font = F_TIT
    ws["A2"] = (f"Corte de estimación · del día {fecha_txt(hasta)}" if desde == hasta else
                f"Corte de estimación · período del {fecha_txt(desde)} al {fecha_txt(hasta)}")
    ws["A2"].font = Font(name="Calibri", bold=True, size=11)
    ws["A3"] = (f"Presupuesto V{p.version_presupuesto} · cantidades cobrables (con tope contractual) · "
                + ("solo lo contractual · " if solo_contractual else "") +
                f"generado {ahora()} por {p.usuario}")
    ws["A3"].font = F_SUB
    enc = ["ITEM", "DESCRIPCION", "UNIDAD", "CANTIDAD", "COSTO UNITARIO", "COSTO TOTAL",
           "CANTIDAD DEL DÍA" if desde == hasta else "CANTIDAD DE ESTE PERÍODO",
           f"MONTO DEL DÍA ({mon})" if desde == hasta else f"MONTO DE ESTE PERÍODO ({mon})"]
    f = 5
    _fila_enc(ws, f, enc)
    for k, w in enumerate((14, 70, 10, 14, 15, 17, 16, 19), start=1):
        ws.column_dimensions[get_column_letter(k)].width = w
    fila = f + 1
    for it in calc.items:
        if not it["vigente"]:
            continue
        r = calc.res[it["id"]]
        if it["tipo"] == "T":
            fill = RELL_T0 if not it["padre"] else RELL_T
            vals = [it["codigo"], it["descripcion"], None, None, None, round(r["total"], 2), None, round(r["m_per"], 2)]
            for k, v in enumerate(vals, start=1):
                c = _poner(ws, fila, k, v, NUM if isinstance(v, float) else None, F_B, fill)
            ws.cell(row=fila, column=2).alignment = Alignment(indent=it["nivel"], wrap_text=True)
        else:
            vals = [it["codigo"], it["descripcion"], it["unidad"], it["cantidad"], it["pu"], it["total"],
                    round(r["cob_per"], 4), round(r["m_per"], 2)]
            for k, v in enumerate(vals, start=1):
                _poner(ws, fila, k, v, NUM if isinstance(v, float) else None)
            ws.cell(row=fila, column=2).alignment = Alignment(indent=it["nivel"], wrap_text=True)
        fila += 1
    t = calc.totales()
    for k, v in enumerate(["TOTAL", None, None, None, None, round(t["total"], 2), None, round(t["m_per"], 2)], start=1):
        _poner(ws, fila, k, v, NUM if isinstance(v, float) else None, F_B, RELL_T0)
    ws.freeze_panes = ws.cell(row=f + 1, column=3)
    _guardar(wb, ruta, guardar)
    return ruta
